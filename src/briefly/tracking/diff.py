"""内容级 Diff：比较同一文件的两个版本，全部用确定性的 Python 完成，不把整份内容交给模型。

- TXT / MD / DOCX / PDF：用现有解析器的结果（Markdown 文本，DOCX 标题和表格已转换）按段落（解析后一段一行）对齐，
  记下所在章节；表格按行比较并附表头；只是换行、空白不同的视为排版噪声，不算内容变化。
  PDF 另按页定位，并标注文本提取可能有误差。
- XLSX：直接读原始文件的单元格（公式按公式文本比较），先按行对齐再比较单元格，插入一行不会让后面所有行都算变化。
- CSV：有可靠主键（唯一、非空、两版大部分重合）时按主键比较；否则按内容对齐，只是行顺序变了时不报变化。
结果带来源版本编号，算法版本变了缓存自动失效，可随时重算。
"""
import csv
import difflib
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from io import BytesIO, StringIO

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker

from .. import config
from ..loaders import decode_text, load_document
from ..materials import MaterialNotFound
from .models import FileSnapshot, SnapshotDiff

ALGO = "diff-v2"  # 输出格式变化时加一，旧缓存自动失效
SIMILAR = 0.5  # 两段相似度达到这个值算“修改”，否则算删除 + 新增
KEY_NAMES = {"id", "编号", "序号", "key", "工号", "订单号", "单号", "code", "代码", "uuid", "名称", "name", "任务"}
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


@dataclass
class Hunk:
    op: str  # added / removed / modified
    location: str  # 章节路径、页码、工作表与单元格、行号或主键
    old: str = ""
    new: str = ""
    context: str = ""  # 表头等辅助信息


@dataclass
class DiffResult:
    file_type: str
    from_snapshot_id: int | None
    to_snapshot_id: int
    hunks: list[Hunk] = field(default_factory=list)
    flags: list[str] = field(default_factory=list)  # format_only / layout_noise / reordered / pdf_text / parse_failed / ...
    notes: list[str] = field(default_factory=list)
    truncated: int = 0  # 超出保留上限、只计数的变化段

    def stats(self) -> dict:
        counts = Counter(h.op for h in self.hunks)
        return {"added": counts["added"], "removed": counts["removed"], "modified": counts["modified"], "truncated": self.truncated}

    def to_dict(self) -> dict:
        return {
            "algo": ALGO,
            "file_type": self.file_type,
            "from_snapshot_id": self.from_snapshot_id,
            "to_snapshot_id": self.to_snapshot_id,
            "hunks": [asdict(h) for h in self.hunks],
            "flags": self.flags,
            "notes": self.notes,
            "truncated": self.truncated,
            "stats": self.stats(),
        }

    @classmethod
    def from_dict(cls, data: dict) -> "DiffResult":
        return cls(
            data["file_type"],
            data["from_snapshot_id"],
            data["to_snapshot_id"],
            [Hunk(**h) for h in data["hunks"]],
            data["flags"],
            data["notes"],
            data["truncated"],
        )

    @property
    def empty(self) -> bool:
        return not self.hunks and not self.truncated


def _norm(text: str) -> str:
    return " ".join(text.split())


def _squash(text: str) -> str:
    return re.sub(r"\s+", "", text)


# ---------- 文本类（txt / md / docx / pdf）
@dataclass
class Unit:
    text: str
    section: str  # 所在章节路径
    line: int  # 在解析文本中的行号（从 1 开始）
    table_header: str = ""  # 表格行所在表格的表头
    para: int = 0  # 第几段（非空行的序号，从 1 开始）：Word、PDF 里用户看到的是段落，不是解析文本的行号


def text_units(text: str, markdown: bool) -> list[Unit]:
    """解析后的文本 → 段落单元（非空行）。Markdown 标题维护章节路径，标题本身也是单元。"""
    units: list[Unit] = []
    path: list[tuple[int, str]] = []
    header, prev_table = "", False
    for n, raw in enumerate(text.splitlines(), start=1):
        line = raw.strip()
        if not line:
            prev_table = False
            continue
        m = _HEADING.match(line) if markdown else None
        if m:
            level = len(m.group(1))
            path = [p for p in path if p[0] < level] + [(level, m.group(2).strip())]
        section = " > ".join(t for _, t in path)
        is_table = line.startswith("|") and line.endswith("|")
        if is_table:
            if re.fullmatch(r"\|[\s:|-]+\|", line):
                continue  # 表头分隔线
            if not prev_table:
                header = line
        units.append(Unit(_norm(line), section, n, header if is_table else "", len(units) + 1))
        prev_table = is_table
    return units


def _location(unit: Unit, kind: str) -> str:
    where = unit.section
    if unit.table_header:
        where = f"{where}（表格）" if where else "表格"
    if where:
        return where
    return f"第 {unit.line} 行" if kind == "txt" else f"第 {unit.para} 段"


def diff_text(old: str | None, new: str, kind: str, markdown: bool) -> tuple[list[Hunk], list[str]]:
    """返回 (变化段, 标记)。old 为 None 表示没有可比较的旧版本：全部内容算新增。"""
    flags: list[str] = []
    new_units = text_units(new, markdown)
    if old is None:
        return _group([Hunk("added", _location(u, kind), new=u.text, context=u.table_header) for u in new_units]), flags
    old_units = text_units(old, markdown)
    if [u.text for u in old_units] == [u.text for u in new_units]:
        return [], flags
    if _squash(old) == _squash(new):
        return [], ["layout_noise"]  # 只是换行、空白不同
    hunks: list[Hunk] = []
    matcher = difflib.SequenceMatcher(None, [u.text for u in old_units], [u.text for u in new_units], autojunk=False)
    for op, i1, i2, j1, j2 in matcher.get_opcodes():
        if op == "equal":
            continue
        olds, news = old_units[i1:i2], new_units[j1:j2]
        if olds and news and _squash("".join(u.text for u in olds)) == _squash("".join(u.text for u in news)):
            if "layout_noise" not in flags:
                flags.append("layout_noise")  # 段落拆分或合并方式变了，文字相同
            continue
        hunks += _pair(olds, news, kind)
    return _group(hunks), flags


def _pair(olds: list[Unit], news: list[Unit], kind: str) -> list[Hunk]:
    """一段被替换的内容：相似的段落配成“修改”，其余算删除或新增。"""
    out: list[Hunk] = []
    used: set[int] = set()
    for o in olds:
        best, best_ratio = None, SIMILAR
        for j, n in enumerate(news):
            if j in used:
                continue
            ratio = difflib.SequenceMatcher(None, o.text, n.text, autojunk=False).ratio()
            if ratio >= best_ratio:
                best, best_ratio = j, ratio
        if best is None:
            out.append(Hunk("removed", _location(o, kind), old=o.text, context=o.table_header))
        else:
            used.add(best)
            n = news[best]
            out.append(Hunk("modified", _location(n, kind), old=o.text, new=n.text, context=n.table_header))
    out += [Hunk("added", _location(n, kind), new=n.text, context=n.table_header) for j, n in enumerate(news) if j not in used]
    return out


def _group(hunks: list[Hunk]) -> list[Hunk]:
    """同一位置连续的新增（或删除）合成一段，减少碎片。"""
    out: list[Hunk] = []
    for h in hunks:
        last = out[-1] if out else None
        if last and h.op == last.op and h.op in ("added", "removed") and h.location == last.location and h.context == last.context:
            if h.op == "added":
                last.new += "\n" + h.new
            else:
                last.old += "\n" + h.old
        else:
            out.append(Hunk(**asdict(h)))
    return out


def pdf_pages(data: bytes) -> list[str]:
    """逐页取文字（去掉空白），只用于给变化段定位页码。"""
    from pypdf import PdfReader

    try:
        return [_squash(p.extract_text() or "") for p in PdfReader(BytesIO(data)).pages]
    except Exception:
        return []


def locate_pages(hunks: list[Hunk], old_pages: list[str], new_pages: list[str]) -> None:
    for h in hunks:
        text, pages = (h.new, new_pages) if h.op != "removed" else (h.old, old_pages)
        probe = _squash(text)[:30]
        if not probe:
            continue
        page = next((i for i, p in enumerate(pages, start=1) if probe in p), None)
        page = page or next((i for i, p in enumerate(pages, start=1) if probe[:12] and probe[:12] in p), None)
        mark = f"第 {page} 页" if page else "页码未能确定"
        h.location = f"{mark} · {h.location}" if h.location else mark


# ---------- 表格类
def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def xlsx_sheets(data: bytes) -> dict[str, list[tuple[str, ...]]]:
    """读原始工作簿：公式按公式文本（以 = 开头）比较，不计算公式结果。"""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), data_only=False, read_only=True)
    sheets = {}
    try:
        for ws in wb.worksheets:
            rows = [tuple(_cell_text(v) for v in row) for row in ws.iter_rows(values_only=True)]
            while rows and not any(rows[-1]):
                rows.pop()
            width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
            sheets[ws.title] = [tuple(r[:width]) + ("",) * (width - len(r[:width])) for r in rows]
    finally:
        wb.close()
    return sheets


def _col(c: int) -> str:
    name = ""
    c += 1
    while c:
        c, rem = divmod(c - 1, 26)
        name = chr(65 + rem) + name
    return name


def _row_text(row: tuple[str, ...], header: tuple[str, ...] | None = None) -> str:
    if header:
        return "，".join(f"{header[i] if i < len(header) and header[i] else _col(i)}={v}" for i, v in enumerate(row) if v)
    return " | ".join(row)


def _is_header(row: tuple[str, ...]) -> bool:
    """第一行像表头：至少一格有字，且都不是数字或公式。"""
    values = [v for v in row if v]
    return bool(values) and not any(v.startswith("=") or re.fullmatch(r"-?[\d.,%]+", v) for v in values)


def diff_rows(
    old: list[tuple[str, ...]], new: list[tuple[str, ...]], where: str, numbered: bool, header_row: bool | None = None
) -> tuple[list[Hunk], list[str]]:
    """按行对齐（内容匹配），被替换的行再比较单元格。numbered：位置用 Excel 行号和列字母。

    header_row：第一行是否表头；None 时按内容判断（两版都至少两行且第一行像表头）。
    """
    flags: list[str] = []
    if header_row is None:
        header_row = len(old) > 1 and len(new) > 1 and _is_header(old[0]) and _is_header(new[0])
    hunks: list[Hunk] = []
    if header_row:
        header = new[0]
        if old[0] != new[0]:
            hunks.append(Hunk("modified", f"{where} 表头", old=" | ".join(old[0]), new=" | ".join(new[0])))
        body_old, body_new, first = old[1:], new[1:], 2
    else:
        header, body_old, body_new, first = (), old, new, 1
    if Counter(body_old) == Counter(body_new):
        if body_old != body_new:
            flags.append("reordered")  # 只是行顺序变了
        return hunks, flags
    matcher = difflib.SequenceMatcher(None, body_old, body_new, autojunk=False)
    for op, i1, i2, j1, j2 in matcher.get_opcodes():
        if op == "equal":
            continue
        olds, news = list(enumerate(body_old[i1:i2], start=i1 + first)), list(enumerate(body_new[j1:j2], start=j1 + first))
        # 被替换的行：按位置配对比较单元格，多出来的算新增或删除
        for (ro, o), (rn, n) in zip(olds, news):
            changes = [
                f"{_col(c) + str(rn) if numbered else (header[c] if c < len(header) and header[c] else _col(c))}："
                f"{o[c] if c < len(o) else ''} → {n[c] if c < len(n) else ''}"
                for c in range(max(len(o), len(n)))
                if (o[c] if c < len(o) else "") != (n[c] if c < len(n) else "")
            ]
            hunks.append(
                Hunk("modified", f"{where} 第 {rn} 行", old=_row_text(o, header), new="；".join(changes), context=" | ".join(header))
            )
        for r, row in olds[len(news):]:
            hunks.append(Hunk("removed", f"{where} 原第 {r} 行", old=_row_text(row, header), context=" | ".join(header)))
        for r, row in news[len(olds):]:
            hunks.append(Hunk("added", f"{where} 第 {r} 行", new=_row_text(row, header), context=" | ".join(header)))
    return hunks, flags


def diff_xlsx(old: bytes | None, new: bytes) -> tuple[list[Hunk], list[str]]:
    new_sheets = xlsx_sheets(new)
    old_sheets = xlsx_sheets(old) if old is not None else {}
    hunks: list[Hunk] = []
    flags: list[str] = []
    for name in old_sheets.keys() - new_sheets.keys():
        hunks.append(Hunk("removed", f"工作表「{name}」", old=f"删除了工作表（{len(old_sheets[name])} 行）"))
    for name, rows in new_sheets.items():
        if name not in old_sheets:
            if old is not None:
                hunks.append(Hunk("added", f"工作表「{name}」", new=f"新增工作表（{len(rows)} 行）"))
            hunks += [Hunk("added", f"「{name}」第 {i} 行", new=_row_text(r, rows[0] if i > 1 else None)) for i, r in enumerate(rows, start=1) if any(r)]
            continue
        sheet_hunks, sheet_flags = diff_rows(old_sheets[name], rows, f"「{name}」", numbered=True)
        hunks += sheet_hunks
        flags += [f for f in sheet_flags if f not in flags]
    return hunks, flags


def _csv_rows(text: str) -> list[tuple[str, ...]]:
    return [tuple(c.strip() for c in row) for row in csv.reader(StringIO(text)) if any(c.strip() for c in row)]


def _find_key(header: tuple[str, ...], old: list[tuple], new: list[tuple]) -> int | None:
    """可靠主键：两版中都唯一、非空，且两版的键大部分重合；优先名字像主键的列。"""
    candidates = sorted(range(len(header)), key=lambda i: header[i].strip().lower() not in KEY_NAMES)
    for i in candidates:
        old_keys = [r[i] if i < len(r) else "" for r in old]
        new_keys = [r[i] if i < len(r) else "" for r in new]
        if not all(old_keys) or not all(new_keys):
            continue
        if len(set(old_keys)) != len(old_keys) or len(set(new_keys)) != len(new_keys):
            continue
        overlap = len(set(old_keys) & set(new_keys)) / max(1, min(len(old_keys), len(new_keys)))
        if overlap >= 0.5:
            return i
    return None


def diff_csv(old: str | None, new: str) -> tuple[list[Hunk], list[str], list[str]]:
    """返回 (变化段, 标记, 说明)。"""
    new_rows = _csv_rows(new)
    if old is None:
        header = new_rows[0] if new_rows else ()
        return [Hunk("added", f"第 {i} 行", new=_row_text(r, header)) for i, r in enumerate(new_rows[1:], start=2)], [], []
    old_rows = _csv_rows(old)
    if not old_rows or not new_rows:
        return diff_rows(old_rows, new_rows, "", numbered=False, header_row=False) + ([],)
    header = new_rows[0]
    key = _find_key(header, old_rows[1:], new_rows[1:]) if old_rows[0] == new_rows[0] else None
    if key is None:
        hunks, flags = diff_rows(old_rows, new_rows, "", numbered=False, header_row=True)
        return hunks, flags, ["没有找到可靠的主键，按行内容对齐比较"]
    name = header[key] or _col(key)
    old_map = {r[key]: r for r in old_rows[1:]}
    new_map = {r[key]: r for r in new_rows[1:]}
    hunks = []
    for k, row in new_map.items():
        if k not in old_map:
            hunks.append(Hunk("added", f"{name}={k}", new=_row_text(row, header)))
        elif old_map[k] != row:
            o = old_map[k]
            changes = [
                f"{header[c] if c < len(header) else _col(c)}：{o[c] if c < len(o) else ''} → {row[c] if c < len(row) else ''}"
                for c in range(max(len(o), len(row)))
                if (o[c] if c < len(o) else "") != (row[c] if c < len(row) else "")
            ]
            hunks.append(Hunk("modified", f"{name}={k}", old=_row_text(o, header), new="；".join(changes)))
    for k, row in old_map.items():
        if k not in new_map:
            hunks.append(Hunk("removed", f"{name}={k}", old=_row_text(row, header)))
    flags = []
    if not hunks and [r[key] for r in old_rows[1:]] != [r[key] for r in new_rows[1:]]:
        flags.append("reordered")
    return hunks, flags, [f"按主键列「{name}」比较"]


# ---------- 对外接口
class Differ:
    """比较同一文件（或可能移动前后的两个文件）的两个版本，结果缓存在 snapshot_diffs。"""

    def __init__(self, svc) -> None:
        self.svc = svc
        self._session = sessionmaker(svc.engine, expire_on_commit=False)

    def text_of(self, snap: FileSnapshot) -> tuple[str | None, bool]:
        """版本的解析文本：优先用材料；材料被删时从原始字节重新解析。返回 (文本, 是否按 Markdown)。"""
        if snap.material_id is not None:
            try:
                m = self.svc.store.get(snap.material_id)
                return m.content, m.file_type in ("md", "doc", "docx", "pdf", "xlsx", "csv") or m.parse_report.get("structure") == "markdown"
            except MaterialNotFound:
                pass
        if snap.parse_error or not self.svc.blobs.has(snap.sha256):
            return None, False
        try:
            doc = load_document(snap.path, self.svc.blobs.get(snap.sha256))
        except ValueError:
            return None, False
        return doc.text.replace("\r\n", "\n").replace("\r", "\n").strip(), bool(doc.markdown)

    def diff(self, old: FileSnapshot | None, new: FileSnapshot, use_cache: bool = True) -> DiffResult:
        if use_cache:
            with self._session() as s:
                cached = s.scalar(
                    select(SnapshotDiff).where(
                        SnapshotDiff.from_snapshot_id == (old.id if old else None),
                        SnapshotDiff.to_snapshot_id == new.id,
                        SnapshotDiff.algo == ALGO,
                    )
                )
            if cached:
                return DiffResult.from_dict(cached.result)
        result = self.compute(old, new)
        with self._session() as s:
            s.add(SnapshotDiff(from_snapshot_id=old.id if old else None, to_snapshot_id=new.id, algo=ALGO, result=result.to_dict()))
            try:
                s.commit()
            except IntegrityError:
                pass  # 并发时别处已存
        return result

    def compute(self, old: FileSnapshot | None, new: FileSnapshot) -> DiffResult:
        kind = new.path.rsplit(".", 1)[-1].lower()
        result = DiffResult(kind, old.id if old else None, new.id)
        if old is not None and old.sha256 == new.sha256:
            return result
        if kind == "xlsx":
            try:
                old_bytes = self.svc.blobs.get(old.sha256) if old else None
                result.hunks, result.flags = diff_xlsx(old_bytes, self.svc.blobs.get(new.sha256))
            except Exception as e:
                result.flags.append("parse_failed")
                result.notes.append(f"工作簿读取失败，无法比较：{type(e).__name__}")
            return self._limit(result)
        if kind == "csv":
            try:
                old_text = decode_text(self.svc.blobs.get(old.sha256)) if old else None
                result.hunks, result.flags, result.notes = diff_csv(old_text, decode_text(self.svc.blobs.get(new.sha256)))
            except (OSError, UnicodeDecodeError, csv.Error) as e:
                result.flags.append("parse_failed")
                result.notes.append(f"CSV 读取失败，无法比较：{type(e).__name__}")
            return self._limit(result)

        new_text, markdown = self.text_of(new)
        old_text = None
        if old is not None:
            old_text, _ = self.text_of(old)
        if new_text is None or (old is not None and old_text is None):
            result.flags.append("parse_failed")
            result.notes.append("有版本解析失败，无法比较内容；原始文件已保存")
            return result
        result.hunks, result.flags = diff_text(old_text, new_text, kind, markdown or kind in ("md", "doc", "docx", "pdf"))
        if not result.hunks and old is not None and kind in ("doc", "docx") and "layout_noise" not in result.flags:
            result.flags.append("format_only")  # 字节变了、正文没变：格式、属性或修订标记等变化
            result.notes.append("文件有变化但正文文字没有变化，可能只改了格式")
        if kind == "pdf":
            result.flags.append("pdf_text")
            result.notes.append("PDF 按提取的文字比较，排版、分栏、OCR 可能造成误差，页码定位仅供参考")
            if result.hunks:
                old_pages = pdf_pages(self.svc.blobs.get(old.sha256)) if old else []
                locate_pages(result.hunks, old_pages, pdf_pages(self.svc.blobs.get(new.sha256)))
        if "layout_noise" in result.flags:
            result.notes.append("有些段落只是换行或空白不同，已忽略")
        return self._limit(result)

    def _limit(self, result: DiffResult) -> DiffResult:
        if len(result.hunks) > config.TRACKING_MAX_HUNKS:
            result.truncated = len(result.hunks) - config.TRACKING_MAX_HUNKS
            result.hunks = result.hunks[: config.TRACKING_MAX_HUNKS]
            result.notes.append(f"变化过多，只保留前 {config.TRACKING_MAX_HUNKS} 处，另有 {result.truncated} 处只计数")
        return result


def format_hunk(h: Hunk, limit: int = 0) -> str:
    """一段变化的文本表示（给用户看，也给模型看）。"""
    def cut(text: str) -> str:
        return text if not limit or len(text) <= limit else text[:limit] + "…"

    where = f"［{h.location}］" if h.location else ""
    if h.op == "added":
        return f"新增{where}：{cut(h.new)}"
    if h.op == "removed":
        return f"删除{where}：{cut(h.old)}"
    return f"修改{where}：{cut(h.old)}\n  改为：{cut(h.new)}"
