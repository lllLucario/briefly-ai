"""表格文件（xlsx、csv）→ Markdown 表格。

xlsx 的每个工作表作为一个章节（# 工作表名）。分块时长表格每块重复表头（见 chunking）。

表头识别（xlsx，见 _header_rows）：
- 表头前只有一格有值的行（表名、“单位:元”）作为说明文字放在表格上方，不当表头——否则表名横向合并后每列都填，所有列都叫表名。
- 多层表头合并成一行，列名写成“上层/下层”（“月工资/标准工资”）；正文中间重复表头的行（分段的子表）去掉。
- 表头和正文都空的列去掉；甘特式排期表连续的日期列压成一列“起～止”（见 _compress_schedule）。

合并单元格（xlsx 与 docx 共用同一规则，见 merged_value）：
- 纵向合并：值向下填到每一行（几个任务共用一个负责人，拆开后每行都要有负责人）。
- 横向合并：值只留在第一格，其余留空（否则“截止”列会出现负责人的名字，误导）。
- 例外：表头行的横向合并每列都填，否则多级表头下的各列看不出属于哪一组。
"""
import csv
import re
from datetime import date, datetime, time
from io import BytesIO, StringIO

from .loaders import LoadedDocument

MAX_HEADER_DEPTH = 4
SCHEDULE_MIN_COLUMNS = 7  # 连续这么多个按序的日期/编号列才算排期
_SCHEDULE_LABEL = re.compile(r"(第?)(\d{1,4})(月|日|号|周|天)?")
_CJK_GAP = re.compile(r"(?<=[一-鿿])\s+(?=[一-鿿])")  # 表头里“标准\n工资”的换行


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel 把整数存成 16.0
    return " ".join(str(value).split())


def _cell(value) -> str:
    return _text(value).replace("|", "\\|")


def merged_value(value, row: int, col: int, first_row: int, first_col: int, header: bool):
    """合并区域内 (row, col) 这一格应显示的值；(first_row, first_col) 是合并区域左上角。"""
    if col == first_col or header:
        return value
    return None


def _display(cell):
    """按单元格的显示格式取值：百分比按 Excel 里看到的样子输出（0.15 → 15%），避免误读。"""
    value = cell.value
    fmt = cell.number_format or ""
    if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" in fmt:
        decimals = len(fmt.split(".")[1].split("%")[0]) if "." in fmt else 0
        return f"{value * 100:.{decimals}f}%"
    return value


def to_markdown_table(rows) -> str:
    cells = [[_cell(v) for v in row] for row in rows]
    cells = [row for row in cells if any(row)]
    if not cells:
        return ""
    width = max(max(i + 1 for i, v in enumerate(row) if v) for row in cells)
    cells = [(row + [""] * width)[:width] for row in cells]
    lines = ["| " + " | ".join(cells[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(row) + " |" for row in cells[1:]]
    return "\n".join(lines)


def load_xlsx(data: bytes) -> LoadedDocument:
    from openpyxl import load_workbook

    # data_only 读 Excel 保存时缓存的计算结果（不用只读模式：要读合并区域和单元格格式）；
    # 另读一份公式，用来发现没有计算结果的公式单元格
    values = load_workbook(BytesIO(data), data_only=True)
    formulas = load_workbook(BytesIO(data), read_only=True)
    parts, missing = [], 0
    try:
        for sheet in values.worksheets:
            width = sheet.max_column
            rows = [[_display(c) for c in row] for row in sheet.iter_rows(min_row=1, min_col=1, max_col=width)]
            formula_rows = formulas[sheet.title].iter_rows(min_row=1, min_col=1, max_col=width, values_only=True)
            for row, formula_row in zip(rows, formula_rows):
                missing += sum(
                    v is None and isinstance(f, str) and f.startswith("=") for v, f in zip(row, formula_row)
                )
            ranges = sheet.merged_cells.ranges
            captions, header = _header_rows(rows, ranges)
            _apply_merges(rows, ranges, header)
            if table := _sheet_table(rows, captions, header):
                parts.append(f"# {sheet.title}\n\n{table}")
    finally:
        values.close()
        formulas.close()
    warnings = (
        [f"有 {missing} 个公式单元格没有计算结果（文件可能不是用 Excel 保存的），这些单元格显示为空"] if missing else []
    )
    report = {"parser": "openpyxl", "structure": "sheets", "warnings": warnings}
    return LoadedDocument("\n\n".join(parts), markdown=True, report=report)


def _filled(row) -> list[int]:
    return [i for i, v in enumerate(row) if _text(v)]


def _header_rows(rows: list[list], ranges) -> tuple[list[str], range]:
    """表头前的说明文字与表头所在的行（合并区域补值之前看原始值）。

    表头首行是第一个至少两格有值的行，它之前的行只作说明。层数看从表头首行开始的合并：
    纵向合并跨几行就是几层；只有横向合并时，下一行的值都落在横向合并的列下面才算第二层（“工时”下的“计划/实际”）。
    """
    nonempty = [i for i, row in enumerate(rows) if _filled(row)]
    if not nonempty:
        return [], range(0)
    first = next((i for i in nonempty if len(_filled(rows[i])) >= 2), None)
    if first is None:  # 单列表：第一行就是表头
        return [], range(nonempty[0], nonempty[0] + 1)
    captions = list(dict.fromkeys(_text(v) for row in rows[:first] for v in row if _text(v)))
    starting = [r for r in ranges if r.min_row - 1 == first]
    depth = max((r.max_row - r.min_row + 1 for r in starting), default=1)
    if depth == 1 and first + 1 < len(rows):
        spans = {c for r in starting if r.max_col > r.min_col for c in range(r.min_col - 1, r.max_col)}
        below = _filled(rows[first + 1])
        if spans and below and set(below) <= spans:
            depth = 2
    return captions, range(first, first + min(depth, MAX_HEADER_DEPTH))


def _apply_merges(rows: list[list], ranges, header: range) -> None:
    """openpyxl 只在合并区域左上角存值，按 merged_value 的规则补到区域内各格。

    正文里横跨半张表以上、又跨多行的合并是说明框（“特别注意事项……”），只留一次，不向下填成多行重复。
    """
    width = len(rows[0]) if rows else 0
    for rng in ranges:
        top, left = rng.min_row - 1, rng.min_col - 1
        value = rows[top][left]
        if top not in header and rng.max_row > rng.min_row and 2 * (rng.max_col - left) >= max(width, 4):
            continue
        for r in range(top, rng.max_row):
            for c in range(left, rng.max_col):
                if (r, c) != (top, left):
                    rows[r][c] = merged_value(value, r, c, top, left, header=r in header)


def _sheet_table(rows: list[list], captions: list[str], header: range) -> str:
    if not header:
        return ""
    width = len(rows[0])
    names = []
    for c in range(width):
        parts = [_CJK_GAP.sub("", _text(rows[h][c])) for h in header]
        parts = [p for i, p in enumerate(parts) if p and p not in parts[:i]]
        names.append("/".join(parts))
    heads = [[_text(v) for v in rows[h]] for h in header]
    body = [[_text(v) for v in row] for row in rows[header.stop:] if _filled(row)]
    body = [row for row in body if not _repeats_header(row, heads)]
    keep = [c for c in range(width) if names[c] or any(row[c] for row in body)]
    names, body = _compress_schedule([names[c] for c in keep], [[row[c] for c in keep] for row in body])
    table = to_markdown_table([names, *body])
    return "\n".join(captions) + "\n\n" + table if captions and table else table


def _repeats_header(row: list[str], heads: list[list[str]]) -> bool:
    """分段的子表在正文里再写一遍表头：有值的格（至少两格）都与某一行表头相同。"""
    cells = {c: v for c, v in enumerate(row) if v}
    return len(cells) >= 2 and any(all(head[c] == v for c, v in cells.items()) for head in heads)


def _ordinal(label: str):
    """排期列的序号与写法：“3”“3月”“第3周”“15日”或日期；都不是返回 None。写法不同的相邻列不算连续。"""
    if match := _SCHEDULE_LABEL.fullmatch(label):
        return match.group(1) + (match.group(3) or ""), int(match.group(2))
    try:
        return "date", date.fromisoformat(label).toordinal()
    except ValueError:
        return None


def _follows(a: str, b: str) -> bool:
    x, y = _ordinal(a), _ordinal(b)
    return x is not None and y is not None and x[0] == y[0] and y[1] == x[1] + 1


def _compress_schedule(names: list[str], body: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """甘特式排期：连续 SCHEDULE_MIN_COLUMNS 个以上列名按序的日期/编号列，正文只有同一种短标记或全空时，
    压成一列，值写有标记的起止（“3～8、12”）。正文是数值的（如 1～12 月的金额）不压，免得丢数。"""
    i = 0
    while i < len(names):
        labels = [n.rsplit("/", 1)[-1] for n in names]
        j = i
        while (j + 1 < len(names) and _follows(labels[j], labels[j + 1])
               and names[j + 1][: -len(labels[j + 1])] == names[i][: -len(labels[i])]):
            j += 1
        marks = {row[c] for row in body for c in range(i, j + 1)} - {""}
        if j - i + 1 < SCHEDULE_MIN_COLUMNS or len(marks) > 1 or any(len(m) > 2 for m in marks):
            i += 1
            continue
        upper = names[i][: -len(labels[i])].rstrip("/") or "排期"
        names[i : j + 1] = [f"{upper}（{labels[i]}～{labels[j]}）"]
        for row in body:
            spans, start = [], None
            for c in range(i, j + 2):
                if c <= j and row[c] and start is None:
                    start = c
                elif (c > j or not row[c]) and start is not None:
                    spans.append(labels[start] if start == c - 1 else f"{labels[start]}～{labels[c - 1]}")
                    start = None
            row[i : j + 1] = ["、".join(spans)]
        i += 1
    return names, body


def load_csv(text: str) -> LoadedDocument:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    table = to_markdown_table(csv.reader(StringIO(text), dialect))
    return LoadedDocument(table, markdown=True, report={"parser": "csv"})
