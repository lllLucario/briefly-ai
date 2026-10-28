from pathlib import Path
import sys
from time import monotonic, sleep

import pytest

from briefly import parse_process
from briefly.parse_process import ParseFailure, load_isolated


def test_real_isolated_text_and_bad_docx():
    doc = load_isolated("说明.txt", "完整保留原文。".encode())
    assert doc.text == "完整保留原文。" and doc.markdown is False
    with pytest.raises(ParseFailure) as error:
        load_isolated("bad.docx", b"broken")
    assert error.value.code == "parse_failed"


def test_real_docx_and_pdf_roundtrip():
    from io import BytesIO
    from docx import Document
    from test_pdf import make_pdf

    document = Document()
    document.add_heading("验收标题", 1)
    document.add_paragraph("独立进程解析，正文保留。")
    stream = BytesIO()
    document.save(stream)
    doc = load_isolated("虚构.docx", stream.getvalue())
    assert "# 验收标题" in doc.text and "正文保留" in doc.text
    pdf = load_isolated("虚构.pdf", make_pdf([[(10, 700, "A sufficiently long original body line for parser testing.")]]))
    assert "original body" in pdf.text


@pytest.mark.parametrize("mode,code", [("exit", "worker_failed"), ("bad", "worker_failed"),
                                      ("empty", "empty_document"), ("error", "parse_failed")])
def test_worker_failures_never_become_success(monkeypatch, mode, code):
    probe = Path(__file__).parent / "fixtures" / "parse_process_probe.py"
    monkeypatch.setattr(parse_process, "_worker_command", lambda folder: [sys.executable, str(probe), str(folder), mode])
    with pytest.raises(ParseFailure) as error:
        load_isolated("bad.doc", b"fake")
    assert error.value.code == code


def test_total_deadline_kills_worker_and_descendants(monkeypatch, tmp_path):
    probe = Path(__file__).parent / "fixtures" / "parse_process_probe.py"
    marker = tmp_path / "leaked.txt"
    folders = []
    def command(folder):
        folders.append(folder)
        return [sys.executable, str(probe), str(folder), "tree", str(marker)]
    monkeypatch.setattr(parse_process, "_worker_command", command)
    started = monotonic()
    with pytest.raises(ParseFailure) as error:
        load_isolated("slow.doc", b"fake", timeout=1)
    assert error.value.code == "parse_timeout"
    assert monotonic() - started < 5
    assert not folders[0].exists()
    sleep(2.2)  # 未清理的后代会在2秒后写marker
    assert not marker.exists()
    # 失败必须释放并发槽位。
    monkeypatch.undo()
    assert load_isolated("ok.txt", b"ok").text == "ok"


def test_queue_wait_counts_against_budget():
    parse_process._slots.acquire()
    parse_process._slots.acquire()
    try:
        with pytest.raises(ParseFailure) as error:
            load_isolated("queue.txt", b"body", timeout=0.02)
        assert error.value.code == "parse_timeout"
    finally:
        parse_process._slots.release()
        parse_process._slots.release()
