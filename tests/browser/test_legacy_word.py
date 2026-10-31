"""旧版Word网页入口：隔离服务、转换替身，无真实模型或Office进程。"""
from io import BytesIO
from pathlib import Path

from docx import Document
import pytest
from playwright.sync_api import expect

from briefly import config, legacy_word

pytestmark=pytest.mark.browser


def test_old_word_upload_and_rich_preview(web_page,monkeypatch,tmp_path):
    monkeypatch.setattr(config,"DOC_PARSER","convert")
    executable=tmp_path/"soffice.com"
    executable.write_bytes(b"fake")
    monkeypatch.setattr(config,"LIBREOFFICE_PATH",str(executable))
    doc=Document()
    doc.add_heading("虚构差旅制度",level=1)
    doc.add_paragraph("住宿上限650元，须事前审批。")
    stream=BytesIO()
    doc.save(stream)
    def fake_conversion(command, **_):
        destination=Path(command[command.index("--outdir")+1])/"source.docx"
        destination.write_bytes(stream.getvalue())
    monkeypatch.setattr(legacy_word,"_run",fake_conversion)
    page=web_page
    assert ".doc" in page.locator("#file").get_attribute("accept").split(",")
    assert ".doc" in page.locator("#todo-upload").get_attribute("accept").split(",")
    page.locator('.tabs [data-tab="materials"]').click()
    page.locator("#open-add").click()
    page.locator("#file").set_input_files({"name":"中文旧制度.doc","mimeType":"application/msword","buffer":legacy_word.OLE_HEADER+b"fixture"})
    expect(page.locator("#add-msg")).to_contain_text("中文旧制度.doc")
    page.locator("#add-dialog [data-close]").first.click()
    expect(page.locator("#material-rows")).to_contain_text("中文旧制度.doc")
    assert page.evaluate("RICH_TYPES.has('doc')") is True


def test_direct_word_upload_uses_local_parser(web_page,monkeypatch,tmp_path):
    from briefly import direct_word
    java,jar=tmp_path/"java.exe",tmp_path/"tika.jar"
    java.write_bytes(b"fake");jar.write_bytes(b"fake")
    monkeypatch.setattr(config,"JAVA_PATH",str(java))
    monkeypatch.setattr(config,"TIKA_APP_JAR",str(jar))
    monkeypatch.setattr(config,"DOC_PARSER","auto")
    payload='<html><head/><body><h1>虚构制度</h1><p>仅正式员工，650元，须事前审批。</p></body></html>'.encode()
    def run(command,**kwargs):
        output=Path(command[command.index("LegacyDocParser")+1])/"0"
        output.mkdir(parents=True)
        (output/"0.xhtml").write_bytes(payload)
    monkeypatch.setattr(direct_word,"_run",run)
    page=web_page
    page.locator('.tabs [data-tab="materials"]').click()
    page.locator("#open-add").click()
    page.locator("#file").set_input_files({"name":"直接解析中文制度.doc","mimeType":"application/msword","buffer":legacy_word.OLE_HEADER+b"fixture"})
    expect(page.locator("#add-msg")).to_contain_text("直接解析中文制度.doc")
    page.locator("#add-dialog [data-close]").first.click()
    expect(page.locator("#material-rows")).to_contain_text("直接解析中文制度.doc")
    materials=page.request.get(page.url.split('#')[0].rstrip('/')+'/materials').json()
    assert any(m['title']=="直接解析中文制度.doc" and m['structure']=="word_styles" for m in materials)
