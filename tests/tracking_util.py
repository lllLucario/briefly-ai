"""文件追踪测试的公共工具：内存库 + 临时版本库，写文件时显式设置 mtime。"""
import os
import time
from io import BytesIO
from pathlib import Path

import pytest

from briefly import config
from briefly.db import create_db_engine
from briefly.materials import MaterialStore
from briefly.tracking.blobs import BlobStore
from briefly.tracking.service import TrackingService
from briefly.weekly.service import WeeklyService


class Clock:
    """测试用的文件修改时间：每次写文件推进一分钟，避免同一秒内的修改被大小 + mtime 筛选漏掉。"""

    def __init__(self) -> None:
        self.t = time.time() - 30 * 24 * 3600

    def next(self) -> float:
        self.t += 60
        return self.t


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "TRACKING_SETTLE_SECONDS", 0)
    monkeypatch.setattr(config, "TRACKING_DIR", str(tmp_path / "store"))
    engine = create_db_engine("sqlite://")
    store = MaterialStore(engine)
    svc = TrackingService(engine, store, BlobStore(tmp_path / "store"))
    work = tmp_path / "work"
    work.mkdir()
    yield Env(svc, store, work, Clock())
    engine.dispose()


class Env:
    def __init__(self, svc: TrackingService, store: MaterialStore, work: Path, clock: Clock) -> None:
        self.svc, self.store, self.work, self.clock = svc, store, work, clock

    def write(self, rel: str, content: str | bytes, mtime: float | None = None) -> Path:
        path = self.work / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(content, str):
            path.write_text(content, encoding="utf-8")
        else:
            path.write_bytes(content)
        t = mtime if mtime is not None else self.clock.next()
        os.utime(path, (t, t))
        return path

    def touch(self, rel: str) -> None:
        t = self.clock.next()
        os.utime(self.work / rel, (t, t))

    def file(self, name: str):
        return next(f for f in self.svc.files() if Path(f.path).name == name)

    def weekly(self, model_factory=lambda: None) -> WeeklyService:
        """共用这个追踪服务的周报服务（测试注入假模型）。"""
        return WeeklyService(self.store, model_factory, tracking=self.svc)

    def set_owners(self, names: list[str]) -> None:
        """你在文档里的名字（周报设置）：归属判断用。"""
        self.weekly().update_settings(owner_names=names)


def docx_bytes(paragraphs: list[str], table: list[tuple[str, ...]] | None = None, headings: dict[int, str] | None = None) -> bytes:
    """生成 docx：paragraphs 按顺序写入；headings 指定在第几段前插入标题（二级）。"""
    from docx import Document

    doc = Document()
    for i, text in enumerate(paragraphs):
        if headings and i in headings:
            doc.add_heading(headings[i], level=2)
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_with_revisions(base: str, inserts: list[tuple[str, str]], last_modified_by: str | None = None) -> bytes:
    """生成带 Word 修订记录（Track Changes）的 docx：inserts 为 [(作者, 插入的文字)]，各成一段。

    python-docx 不支持修订，直接改 document.xml；last_modified_by 写进文档属性。
    """
    import zipfile

    from docx import Document

    doc = Document()
    doc.add_paragraph(base)
    for i, (_, text) in enumerate(inserts):
        doc.add_paragraph(f"@@INS{i}@@")
    if last_modified_by:
        doc.core_properties.last_modified_by = last_modified_by
    buf = BytesIO()
    doc.save(buf)
    src = zipfile.ZipFile(BytesIO(buf.getvalue()))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                for i, (author, text) in enumerate(inserts):
                    run = f"<w:r><w:t>@@INS{i}@@</w:t></w:r>"
                    ins = f'<w:ins w:id="{i + 1}" w:author="{author}" w:date="2026-10-14T10:00:00Z"><w:r><w:t>{text}</w:t></w:r></w:ins>'
                    assert run in xml
                    xml = xml.replace(run, ins)
                data = xml.encode("utf-8")
            z.writestr(item, data)
    return out.getvalue()


def xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_bytes(pages: list[list[str]]) -> bytes:
    """生成只含 ASCII 文字的最小 PDF：每页若干行（Helvetica 12 号），用于测试 PDF 文本比较。"""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for lines in pages:
        ops = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
        for line in lines:
            safe = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            ops.append(f"({safe}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops)
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_id = len(objects)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            "/Resources << /Font << /F1 3 0 R >> >> >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out
