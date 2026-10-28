from io import BytesIO
from pathlib import Path
import subprocess

import pytest
from docx import Document
from fastapi.testclient import TestClient

from briefly import config, legacy_word
from briefly.api import create_app
from briefly.db import create_db_engine
from briefly.loaders import file_type, load_document
from briefly.materials import MaterialStore


def docx_bytes():
    doc=Document()
    doc.add_heading("虚构住宿制度",level=1)
    doc.add_paragraph("上海住宿上限650元，须事前审批。")
    table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text="城市"
    table.cell(0,1).text="额度"
    table.cell(1,0).text="杭州"
    table.cell(1,1).text="480元"
    output=BytesIO()
    doc.save(output)
    return output.getvalue()


@pytest.fixture
def converted(monkeypatch,tmp_path):
    monkeypatch.setattr(config,"DOC_PARSER","convert")
    executable=tmp_path/"中文 程序"/"soffice.com"
    executable.parent.mkdir()
    executable.write_bytes(b"fake")
    monkeypatch.setattr(config,"LIBREOFFICE_PATH",str(executable))
    commands=[]
    def run(command):
        commands.append(command)
        source=Path(command[-1])
        assert source.read_bytes().startswith(legacy_word.OLE_HEADER)
        profile=next(c.split("=",1)[1] for c in command if c.startswith("-env:UserInstallation="))
        assert profile.startswith("file:///")
        destination=Path(command[command.index("--outdir")+1])/"source.docx"
        destination.write_bytes(docx_bytes())
    monkeypatch.setattr(legacy_word,"_run",run)
    return commands


def test_convert_then_preserve_headings_tables_and_original_type(converted):
    doc=load_document("中文目录/制度.DOC",legacy_word.OLE_HEADER+b"fixture")
    assert file_type("制度.DOC")=="doc"
    assert "# 虚构住宿制度" in doc.text
    assert "650元" in doc.text and "| 杭州 | 480元 |" in doc.text
    assert doc.markdown is True
    assert doc.report["original_format"]=="doc" and doc.report["warnings"]
    assert not Path(converted[0][-1]).exists()  # 临时文件已清理


def test_missing_runtime_and_misconfigured_path(monkeypatch):
    monkeypatch.setattr(config,"DOC_PARSER","convert")
    monkeypatch.setattr(config,"LIBREOFFICE_PATH","")
    monkeypatch.setattr(legacy_word.shutil,"which",lambda _:None)
    monkeypatch.setattr(Path,"is_file",lambda _:False)
    with pytest.raises(ValueError,match="需要服务端安装"):
        load_document("a.doc",legacy_word.OLE_HEADER+b"fixture")
    monkeypatch.setattr(config,"LIBREOFFICE_PATH","missing-soffice")
    with pytest.raises(ValueError,match="指定的转换程序不存在"):
        legacy_word.converter_path()


@pytest.mark.parametrize("payload",[b"",b"wrong",b"PK\x03\x04wrong-docx-renamed"])
def test_invalid_or_renamed_document_is_not_success(payload):
    with pytest.raises(ValueError,match="仅更改后缀"):
        load_document("a.doc",payload)


def test_no_output_or_invalid_output_is_not_success(converted,monkeypatch):
    monkeypatch.setattr(legacy_word,"_run",lambda _:None)
    with pytest.raises(ValueError,match="未获得有效转换结果"):
        load_document("a.doc",legacy_word.OLE_HEADER)
    def invalid(command):
        (Path(command[command.index("--outdir")+1])/"source.docx").write_bytes(b"bad")
    monkeypatch.setattr(legacy_word,"_run",invalid)
    with pytest.raises(ValueError,match="不是有效"):
        load_document("a.doc",legacy_word.OLE_HEADER)


def test_busy_slots_fail_clearly(converted,monkeypatch):
    class Busy:
        def acquire(self,**kwargs):
            return False
    monkeypatch.setattr(legacy_word,"_slots",Busy())
    with pytest.raises(ValueError,match="繁忙"):
        load_document("a.doc",legacy_word.OLE_HEADER)


def test_timeout_and_nonzero_exit(monkeypatch):
    class Process:
        pid=12345
        returncode=1
        def communicate(self,**kwargs):
            raise subprocess.TimeoutExpired("fake",1)
        def wait(self,**kwargs):
            pass
    monkeypatch.setattr(legacy_word.subprocess,"Popen",lambda *a,**kw:Process())
    monkeypatch.setattr(legacy_word.subprocess,"run",lambda *a,**kw:None)
    monkeypatch.setattr(legacy_word.os,"killpg",lambda *args:None,raising=False)
    with pytest.raises(ValueError,match="超时"):
        legacy_word._run(["fake"])
    monkeypatch.setattr(Process,"communicate",lambda *a,**kw:None)
    with pytest.raises(ValueError,match="转换失败"):
        legacy_word._run(["fake"])


def test_upload_old_word_through_existing_api(converted):
    engine=create_db_engine("sqlite://")
    try:
        with TestClient(create_app(MaterialStore(engine))) as client:
            result=client.post("/materials/upload",files={"file":("中文制度.doc",legacy_word.OLE_HEADER+b"fixture")})
            assert result.status_code==201
            material=result.json()["material"]
            assert material["file_type"]=="doc" and material["structure"]=="docx_styles"
            detail=client.get(f"/materials/{material['id']}").json()
            assert "650元" in detail["content"]
    finally:
        engine.dispose()
