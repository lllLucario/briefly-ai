from io import BytesIO

import pytest
from langchain_core.messages import AIMessage, HumanMessage, ToolMessage

from briefly import config
from briefly.agent import Agent
from briefly.context import Context
from briefly.db import create_db_engine
from briefly.materials import MaterialStore
from fakes import FunctionModel, fake, structured, text_msg, tool_msg


@pytest.fixture
def store():
    engine = create_db_engine("sqlite://")
    yield MaterialStore(engine)
    engine.dispose()


def run_tool(store: MaterialStore, name: str, args: dict, material_ids: list[int] | None = None) -> ToolMessage:
    """让假模型调用一次工具，返回工具结果消息。"""
    model = fake(tool_msg(name, args), text_msg("ok"))
    agent = Agent(model, tools=[], store=store)
    agent.graph.invoke(
        {"messages": [HumanMessage("x")]},
        config={"configurable": {"thread_id": "t"}},
        context=Context(material_ids=material_ids),
    )
    return model.calls[1][-1]


def docx_bytes(*paragraphs: str) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---- 服务层 ----


def test_add_text_and_get(store):
    material, created = store.add_text("周报素材", "  本周完成了接口联调。  ")
    assert created
    got = store.get(material.id)
    assert got.content == "本周完成了接口联调。"
    assert got.char_count == 10 and got.source == "paste" and got.file_type is None


def test_blank_title_falls_back_to_content(store):
    material, _ = store.add_text("", "会议纪要：讨论了上线计划")
    assert material.title == "会议纪要：讨论了上线计划"


def test_empty_text_rejected(store):
    with pytest.raises(ValueError, match="为空"):
        store.add_text("t", "   \n ")


def test_duplicate_content_returns_existing(store):
    first, _ = store.add_text("a", "相同内容")
    second, created = store.add_text("b", "相同内容")
    assert not created and second.id == first.id
    assert len(store.list_materials()) == 1


def test_add_txt_file_utf8_and_gbk(store):
    utf8, _ = store.add_file("a.txt", "完成登录模块".encode("utf-8"))
    gbk, _ = store.add_file("b.md", "修复支付缺陷".encode("gbk"))
    assert store.get(utf8.id).content == "完成登录模块"
    assert store.get(gbk.id).content == "修复支付缺陷"
    assert gbk.file_type == "md" and gbk.source == "file" and gbk.title == "b.md"


def test_windows_line_endings_are_normalized(store):
    material, _ = store.add_file("a.txt", "第一行\r\n第二行\r\n".encode("utf-8"))
    assert material.content == "第一行\n第二行" and material.char_count == 7
    _, created = store.add_text("b", "第一行\n第二行")
    assert not created  # 换行不同不影响去重


def test_add_docx_file(store):
    material, _ = store.add_file("纪要.docx", docx_bytes("第一段", "第二段", "# Role 正文里的示例"))
    assert material.content == "第一段\n\n第二段\n\n\\# Role 正文里的示例"
    assert material.parse_report["parser"] == "docling"
    assert store.sections(material.id) == []  # 正文里的 # 已转义，不当标题


def test_docx_table_and_special_characters_are_kept(store):
    from docx import Document

    doc = Document()
    doc.add_paragraph("变量 file_name 和 A&B")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table._cells, ["任务", "负责人", "支付|改造", "李经理"]):
        cell.text = text
    buf = BytesIO()
    doc.save(buf)
    material, _ = store.add_file("t.docx", buf.getvalue())
    assert "变量 file_name 和 A&B" in material.content  # 不转义成 file\_name、A&amp;B
    assert "| 任务 | 负责人 |" in material.content
    assert "支付\\|改造" in material.content


def test_corrupted_file_gives_readable_error(store):
    with pytest.raises(ValueError, match="文件无法解析"):
        store.add_file("坏文件.docx", b"not a zip file")


def test_pdf_without_text_rejected(store):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = BytesIO()
    writer.write(buf)
    with pytest.raises(ValueError, match="扫描版"):
        store.add_file("scan.pdf", buf.getvalue())


def test_unsupported_and_oversized_files_rejected(store, monkeypatch):
    with pytest.raises(ValueError, match="不支持"):
        store.add_file("a.pptx", b"x")
    monkeypatch.setattr(config, "MAX_UPLOAD_MB", 0)
    with pytest.raises(ValueError, match="上限"):
        store.add_file("a.txt", b"x")


def test_delete(store):
    material, _ = store.add_text("t", "内容")
    store.delete(material.id)
    assert store.list_materials() == []
    with pytest.raises(ValueError, match="不存在"):
        store.get(material.id)


def test_list_materials_filters_by_ids(store):
    a, _ = store.add_text("a", "甲")
    store.add_text("b", "乙")
    assert [m.id for m in store.list_materials([a.id])] == [a.id]


def test_sections_and_chunks_are_stored_on_import(store, monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE", 30)
    monkeypatch.setattr(config, "CHUNK_OVERLAP", 0)
    material, _ = store.add_text("周报", "# 周报\n## 周一\n完成登录模块联调工作。\n## 周二\n评审订单分页方案。")
    assert [(s.index, s.level, s.path) for s in store.sections(material.id)] == [
        (1, 1, "周报"),
        (2, 2, "周报 > 周一"),
        (3, 2, "周报 > 周二"),
    ]
    chunks = store.chunks(material.id)
    assert [c.index for c in chunks] == list(range(1, len(chunks) + 1))
    assert all(material.content[c.start : c.end].strip() == c.content for c in chunks)
    assert chunks[-1].section == "周报 > 周二"
    assert store.chunk_counts([material.id]) == {material.id: len(chunks)}


def test_parse_report_records_structure_source(store):
    md, _ = store.add_file("a.md", "# 标题\n# Role 是正文吗？在 md 里算标题".encode())
    txt, _ = store.add_file("b.txt", "一、背景\n# Role\n正文".encode())
    plain, _ = store.add_file("c.txt", "没有任何标题的一段话".encode())
    pasted, _ = store.add_text("粘贴", "第一章 总则\n内容")
    assert md.parse_report["structure"] == "markdown"
    assert txt.parse_report["structure"] == "rules"
    assert [s.title for s in store.sections(txt.id)] == ["一、背景"]
    assert plain.parse_report["structure"] == "none"
    assert pasted.parse_report["structure"] == "rules"


def test_delete_removes_sections_and_chunks(store):
    from sqlalchemy import func, select

    from briefly.materials import MaterialChunk, MaterialSection

    material, _ = store.add_text("t", "# 标题\n内容")
    store.delete(material.id)
    with store._session() as s:
        assert s.scalar(select(func.count()).select_from(MaterialChunk)) == 0
        assert s.scalar(select(func.count()).select_from(MaterialSection)) == 0


def test_docx_heading_styles_become_sections(store):
    from docx import Document

    doc = Document()
    doc.add_heading("项目周报", level=1)
    doc.add_heading("本周完成", level=2)
    doc.add_paragraph("完成登录模块")
    buf = BytesIO()
    doc.save(buf)
    material, _ = store.add_file("周报.docx", buf.getvalue())
    assert [s.path for s in store.sections(material.id)] == ["项目周报", "项目周报 > 本周完成"]
    assert material.parse_report["structure"] == "docx_styles"


def test_search_ranks_by_hit_count(store):
    a, _ = store.add_text("a", "接口联调完成。")
    b, _ = store.add_text("b", "接口文档更新，接口测试通过。")
    hits = store.search("接口")
    assert [h.material_id for h in hits] == [b.id, a.id]
    assert hits[0].hits == 2 and hits[0].chunk == 1


def test_search_multiple_keywords_and_case(store):
    store.add_text("a", "Release 计划已确认")
    store.add_text("b", "无关内容")
    assert len(store.search("release 预算")) == 1


def test_search_snippet_is_trimmed(store):
    store.add_text("a", "甲" * 300 + "关键" + "乙" * 300)
    snippet = store.search("关键")[0].snippet
    assert "关键" in snippet and snippet.startswith("…") and snippet.endswith("…")
    assert len(snippet) < 300


# ---- 工具 ----


def test_list_materials_tool(store):
    store.add_text("周报素材", "内容")
    result = run_tool(store, "list_materials", {})
    assert "材料 1《周报素材》" in result.content and "没有识别出章节" in result.content
    assert "块" not in result.content and "字" not in result.content.split("》", 1)[1]  # 块数、字数会被当成文档内容


def test_read_material_tool_wraps_content(store):
    material, _ = store.add_text("t", "本周完成接口联调")
    result = run_tool(store, "read_material", {"material_id": material.id})
    assert "第 1/1 块" in result.content
    assert "【材料开始】\n本周完成接口联调\n【材料结束】" in result.content


def test_read_material_chunk_out_of_range(store):
    material, _ = store.add_text("t", "内容")
    result = run_tool(store, "read_material", {"material_id": material.id, "chunk": 2})
    assert "工具执行失败" in result.content and "共 1 块" in result.content


def test_search_materials_tool(store):
    store.add_text("t", "接口联调完成")
    result = run_tool(store, "search_materials", {"keywords": "联调"})
    assert "第 1 块，命中 1 次" in result.content
    assert "没有找到" in run_tool(store, "search_materials", {"keywords": "预算"}).content


def test_tools_respect_selected_scope(store):
    a, _ = store.add_text("a", "接口联调")
    b, _ = store.add_text("b", "接口测试")
    listed = run_tool(store, "list_materials", {}, material_ids=[a.id])
    assert "《a》" in listed.content and "《b》" not in listed.content
    denied = run_tool(store, "read_material", {"material_id": b.id}, material_ids=[a.id])
    assert "不在用户本次选择的范围内" in denied.content
    searched = run_tool(store, "search_materials", {"keywords": "接口"}, material_ids=[a.id])
    assert "《a》" in searched.content and "《b》" not in searched.content
    denied = run_tool(store, "search_materials", {"keywords": "接口", "material_ids": [b.id]}, material_ids=[a.id])
    assert "不在用户本次选择的范围内" in denied.content


WEEKLY = "# 周报\n## 周一\n完成登录联调。\n## 周三\n### 支付\n定位回调超时。\n### 订单\n评审分页方案。\n## 周五\n确认上线时间。"


def test_get_outline_lists_sections_without_internal_numbers(store):
    material, _ = store.add_text("周报", WEEKLY)
    result = run_tool(store, "get_outline", {"material_id": material.id}).content
    assert result.startswith("【材料 1《周报》的章节大纲") and "个章节" not in result.splitlines()[0]
    assert "\n1. 周报\n" in result and "块" not in result and "字）" not in result  # 不写块号、字数
    assert "\n    4. 支付\n" in result  # 三级标题缩进两层
    assert "【材料开始】" in result  # 标题来自材料，同样当作材料


def test_get_outline_without_sections_points_to_search(store):
    material, _ = store.add_text("t", "没有任何标题的一段话。")
    result = run_tool(store, "get_outline", {"material_id": material.id}).content
    assert "没有识别出章节" in result and "search_materials" in result


def test_read_material_by_section_returns_only_that_section(store):
    material, _ = store.add_text("周报", WEEKLY)
    result = run_tool(store, "read_material", {"material_id": material.id, "section": 3}).content
    assert "章节 3：周报 > 周三" in result
    assert "定位回调超时" in result and "评审分页方案" in result  # 含子章节
    assert "完成登录联调" not in result and "确认上线时间" not in result


def test_long_section_returns_first_chunks_and_where_to_continue(store, monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE", 40)
    monkeypatch.setattr(config, "CHUNK_OVERLAP", 0)
    monkeypatch.setattr(config, "SECTION_READ_CHARS", 90)
    body = "\n".join(f"第{i}条记录的内容比较长一些。" for i in range(20))
    material, _ = store.add_text("长", f"# 前言\n简介。\n# 正文\n{body}")
    chunks = store.chunks(material.id)
    result = run_tool(store, "read_material", {"material_id": material.id, "section": 2}).content
    assert "较长，本次返回" in result
    assert f"读到第 {len(chunks)} 块" in result
    assert "第0条记录" in result and "第19条记录" not in result


def test_read_material_rejects_both_chunk_and_section(store):
    material, _ = store.add_text("周报", WEEKLY)
    result = run_tool(store, "read_material", {"material_id": material.id, "chunk": 1, "section": 1}).content
    assert "只能填一个" in result


def test_read_material_section_out_of_range(store):
    material, _ = store.add_text("周报", WEEKLY)
    result = run_tool(store, "read_material", {"material_id": material.id, "section": 99}).content
    assert "章节编号超出范围" in result and "get_outline" in result


def test_chunk_read_and_search_show_section_path(store, monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE", 20)
    monkeypatch.setattr(config, "CHUNK_OVERLAP", 0)
    material, _ = store.add_text("周报", WEEKLY)
    searched = run_tool(store, "search_materials", {"keywords": "回调"}).content
    assert "（周报 > 周三 > 支付），命中 1 次" in searched
    chunk = next(c for c in store.chunks(material.id) if "回调" in c.content)
    read = run_tool(store, "read_material", {"material_id": material.id, "chunk": chunk.index}).content
    assert f"块（{chunk.section}）】" in read


def test_get_outline_respects_selected_scope(store):
    a, _ = store.add_text("a", "# 甲\n内容")
    b, _ = store.add_text("b", "# 乙\n内容")
    result = run_tool(store, "get_outline", {"material_id": b.id}, material_ids=[a.id]).content
    assert "不在用户本次选择的范围内" in result


def run_summarize(
    store, args: dict, material_ids=None, extract=None, digest: str = "- 支付回调耗时下降〔#1〕"
) -> tuple[str, FunctionModel]:
    """智能体调用 summarize_materials；工具内部的提取、汇总也由同一个假模型按系统提示词应答。"""

    def respond(messages):
        system = messages[0].content
        if "你负责从一段材料中提取要点" in system:
            if extract:
                return extract(messages)
            return structured("Extraction", items=[{"category": "done", "point": "耗时下降", "quote": "耗时降到 200 毫秒。"}])
        if "你负责把从材料中逐块提取" in system:
            return AIMessage(digest)
        if isinstance(messages[-1], ToolMessage):
            return AIMessage("整理好的总结")
        return tool_msg("summarize_materials", args)

    model = FunctionModel(respond=respond)
    agent = Agent(model, tools=[], store=store)
    agent.graph.invoke(
        {"messages": [HumanMessage("总结这份材料")]},
        config={"configurable": {"thread_id": agent.thread_id}},
        context=Context(material_ids=material_ids),
    )
    tool_result = next(m for m in agent.history if isinstance(m, ToolMessage))
    return tool_result.content, model


def test_summarize_materials_returns_verified_digest(store):
    store.add_text("周报", "# 周三\n支付回调耗时降到 200 毫秒。")
    result, model = run_summarize(store, {"focus": "支付"})
    assert "材料已全部处理" in result and "〔〕内为出处" in result
    assert "【材料开始】\n- 支付回调耗时下降〔《周报》周三〕\n【材料结束】" in result  # 出处用标题和章节，不用块号
    extract_call = next(c for c in model.calls if "你负责从一段材料中提取要点" in c[0].content)
    assert "简述（用户说明的重点，不是筛选条件）：支付" in extract_call[-1].content
    digest_call = next(c for c in model.calls if "你负责把从材料中逐块提取" in c[0].content)
    assert digest_call[-1].content.startswith("关注点：支付")


def test_summarize_materials_respects_scope_and_reports_failures(store):
    a, _ = store.add_text("a", "耗时降到 200 毫秒。")
    b, _ = store.add_text("b", "另一份材料。")
    denied, _ = run_summarize(store, {"material_ids": [b.id]}, material_ids=[a.id])
    assert "不在用户本次选择的范围内" in denied

    def extract(messages):
        if "《b》" in messages[-1].content:
            raise ConnectionError("网络错误")
        return structured("Extraction", items=[{"category": "other", "point": "耗时", "quote": "耗时降到 200 毫秒。"}])

    result, _ = run_summarize(store, {}, extract=extract)
    assert "《b》有部分内容处理失败，未纳入" in result and "块" not in result


def test_summarize_materials_appends_points_the_digest_left_out(store):
    store.add_text("周报", "定位超时原因。耗时降到 200 毫秒。")

    def extract(messages):
        return structured(
            "Extraction",
            items=[
                {"category": "done", "point": "定位超时原因", "quote": "定位超时原因。"},
                {"category": "done", "point": "耗时下降", "quote": "耗时降到 200 毫秒。"},
            ],
        )

    result, _ = run_summarize(store, {}, extract=extract, digest="- 耗时下降〔#2〕")
    assert "汇总时未写入的要点：\n- 定位超时原因〔《周报》〕" in result


def test_summarize_materials_with_nothing_relevant(store):
    store.add_text("a", "整理需求池。")
    result, _ = run_summarize(store, {"focus": "预算"}, extract=lambda m: structured("Extraction", items=[]))
    assert result.startswith("没有提取到与“预算”相关的内容")


def test_reading_beyond_turn_budget_suggests_summarize(store, monkeypatch):
    monkeypatch.setattr(config, "TURN_TOOL_TOKENS", 350)  # 第一次读完后剩余不足 300，第二次不再执行
    material, _ = store.add_text("长", "第一块的材料原文内容比较长。" * 5)
    model = fake(
        tool_msg("read_material", {"material_id": material.id, "chunk": 1}, id="r1"),
        tool_msg("read_material", {"material_id": material.id, "chunk": 1}, id="r2"),
        text_msg("ok"),
        tool_msg("read_material", {"material_id": material.id, "chunk": 1}, id="r3"),
        text_msg("ok"),
    )
    agent = Agent(model, tools=[], store=store)
    agent.chat("读一下")
    first, second = [m.content for m in agent.history if isinstance(m, ToolMessage)]
    assert "【材料开始】" in first
    assert "达到单轮预算" in second and "summarize_materials" in second
    agent.chat("再读一次")  # 新一轮重新计算
    assert "【材料开始】" in [m.content for m in agent.history if isinstance(m, ToolMessage)][2]


def test_runtime_is_hidden_from_model_schema(store):
    from briefly.tools.materials import material_tools

    for t in material_tools(store):
        assert "runtime" not in t.tool_call_schema.model_json_schema().get("properties", {})


# ---- 智能体 ----


def test_short_materials_are_inlined(store):
    a, _ = store.add_text("素材A", "完成登录模块")
    model = fake(text_msg("ok"))
    Agent(model, tools=[], store=store).chat("写周报", skill="weekly", material_ids=[a.id])
    message = model.calls[0][-1].content
    assert message.startswith("写周报")
    assert "【材料 1《素材A》】\n【材料开始】\n完成登录模块\n【材料结束】" in message


def test_long_materials_are_listed_not_inlined(store, monkeypatch):
    monkeypatch.setattr(config, "INLINE_MATERIAL_TOKENS", 5)
    a, _ = store.add_text("素材A", "完成登录模块和支付模块")
    model = fake(text_msg("ok"))
    Agent(model, tools=[], store=store).chat("写周报", material_ids=[a.id])
    message = model.calls[0][-1].content
    assert "【已选材料】" in message and "材料 1《素材A》" in message
    assert "完成登录模块" not in message


def test_selecting_missing_material_raises(store):
    with pytest.raises(ValueError, match="不存在"):
        Agent(fake(), tools=[], store=store).chat("x", material_ids=[99])


def test_search_then_read_then_answer(store, monkeypatch):
    monkeypatch.setattr(config, "INLINE_MATERIAL_TOKENS", 0)
    a, _ = store.add_text("周报素材", "周一：接口联调。周三：修复支付缺陷。")
    model = fake(
        tool_msg("search_materials", {"keywords": "支付"}, id="s"),
        tool_msg("read_material", {"material_id": a.id, "chunk": 1}, id="r"),
        text_msg("本周修复了支付缺陷。"),
    )
    reply = Agent(model, tools=[], store=store).chat("写周报", material_ids=[a.id])
    assert reply == "本周修复了支付缺陷。"
    assert "命中 1 次" in model.calls[1][-1].content
    assert "周三：修复支付缺陷" in model.calls[2][-1].content
