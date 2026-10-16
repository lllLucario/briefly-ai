"""文件追踪 Step 2：内容级 Diff（TXT、MD、DOCX、PDF、XLSX、CSV）。确定性计算，不调用模型。"""
from io import BytesIO

from briefly import config
from briefly.tracking import cli as track_cli
from briefly.tracking.diff import Differ, diff_csv, diff_text, diff_xlsx, format_hunk
from tracking_util import docx_bytes, env, pdf_bytes, xlsx_bytes  # noqa: F401


def ops(hunks):
    return [(h.op, h.location) for h in hunks]


def track(env, name, content):
    """追踪一个文件，返回这个文件。"""
    env.write(name, content)
    env.svc.add_paths([str(env.work / name)])
    return env.file(name.split("/")[-1])


def update(env, f, name, content):
    env.write(name, content)
    env.svc.scan()
    snaps = env.svc.snapshots(f.id)
    return Differ(env.svc).diff(snaps[-2], snaps[-1])


def test_markdown_diff_keeps_sections():
    old = "# 方案\n## 背景\n原有背景说明。\n## 架构\n使用单体服务。\n## 计划\n十月上线。"
    new = "# 方案\n## 背景\n原有背景说明。\n## 架构\n使用单体服务，增加异步队列。\n新增缓存层。\n## 计划\n十月上线。\n## 风险\n压测未完成。"
    hunks, flags = diff_text(old, new, "md", markdown=True)
    assert ops(hunks) == [
        ("modified", "方案 > 架构"),
        ("added", "方案 > 架构"),
        ("added", "方案 > 风险"),
    ]
    assert hunks[0].old == "使用单体服务。" and "异步队列" in hunks[0].new
    assert hunks[2].new == "## 风险\n压测未完成。"  # 连续新增合成一段
    assert flags == []


def test_txt_line_diff_uses_line_numbers():
    hunks, _ = diff_text("第一行\n第二行\n第三行", "第一行\n第二行改了一些字\n第三行\n第四行", "txt", markdown=False)
    assert ops(hunks) == [("modified", "第 2 行"), ("added", "第 4 行")]


def test_documents_without_headings_use_paragraph_numbers():
    hunks, _ = diff_text("第一段\n\n第二段内容", "第一段\n\n第二段内容，补充说明\n\n第三段", "docx", markdown=True)
    assert ops(hunks) == [("modified", "第 2 段"), ("added", "第 3 段")]  # 不是解析文本的行号（中间有空行）


def test_whitespace_and_reflow_are_layout_noise():
    hunks, flags = diff_text("系统采用异步队列\n处理回调。", "系统采用异步队列处理回调。", "pdf", markdown=True)
    assert hunks == [] and flags == ["layout_noise"]
    hunks, flags = diff_text("甲\n乙  丙", "甲\n乙 丙", "txt", markdown=False)
    assert hunks == [] and flags == []


def test_first_version_is_all_added():
    hunks, _ = diff_text(None, "# 新文件\n第一段\n第二段", "md", markdown=True)
    assert [h.op for h in hunks] == ["added"] and "第二段" in hunks[0].new


def test_docx_paragraph_and_table_changes(env):
    old = docx_bytes(["项目进展说明", "接口联调进行中"], table=[("任务", "状态"), ("压测", "未开始")], headings={1: "进度"})
    f = track(env, "周报.docx", old)
    new = docx_bytes(
        ["项目进展说明", "接口联调已完成，进入测试", "新增：支付回调超时排查"],
        table=[("任务", "状态"), ("压测", "进行中"), ("回归", "未开始")],
        headings={1: "进度"},
    )
    result = update(env, f, "周报.docx", new)
    texts = [format_hunk(h) for h in result.hunks]
    assert any(h.op == "modified" and "联调已完成" in h.new and h.location == "进度" for h in result.hunks)
    assert any(h.op == "added" and "支付回调超时" in h.new for h in result.hunks)
    table = [h for h in result.hunks if h.context]
    assert table and all("任务" in h.context for h in table)  # 表格变化附表头
    assert any("进行中" in h.new for h in table) and any("回归" in h.new for h in table)
    assert result.from_snapshot_id and result.to_snapshot_id, texts


def test_docx_format_only_change(env):
    from docx import Document
    from docx.shared import Pt

    def build(size):
        doc = Document()
        doc.add_paragraph().add_run("正文内容不变").font.size = Pt(size)
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    f = track(env, "a.docx", build(11))
    result = update(env, f, "a.docx", build(16))
    assert result.hunks == [] and "format_only" in result.flags


def test_pdf_text_change_is_located_by_page_and_marked_uncertain(env):
    f = track(env, "spec.pdf", pdf_bytes([["Project plan", "The system uses a queue."], ["Owner is Li.", "Deadline is October."]]))
    result = update(
        env, f, "spec.pdf", pdf_bytes([["Project plan", "The system uses a queue."], ["Owner is Wang.", "Deadline is October."]])
    )
    assert [h.op for h in result.hunks] == ["modified"]
    assert result.hunks[0].location.startswith("第 2 页") and "Wang" in result.hunks[0].new
    assert "pdf_text" in result.flags


def test_pdf_reflow_is_not_a_change(env):
    f = track(env, "a.pdf", pdf_bytes([["The system uses a queue and", "a cache for callbacks."]]))
    result = update(env, f, "a.pdf", pdf_bytes([["The system uses a queue and a", "cache for callbacks."]]))
    assert result.hunks == []


def test_xlsx_cell_row_sheet_and_formula_changes():
    old = xlsx_bytes({"进度": [["任务", "负责人", "进度"], ["联调", "李强", 50], ["压测", "赵琳", 0]], "汇总": [["合计", "=SUM(1,2)"]]})
    new = xlsx_bytes(
        {
            "进度": [["任务", "负责人", "进度"], ["联调", "李强", 100], ["回调排查", "我", 30], ["压测", "赵琳", 0]],
            "汇总": [["合计", "=SUM(1,2,3)"]],
            "风险": [["事项"], ["数据库超时"]],
        }
    )
    hunks, flags = diff_xlsx(old, new)
    text = "\n".join(format_hunk(h) for h in hunks)
    assert "C2：50 → 100" in text
    assert any(h.op == "added" and "回调排查" in h.new for h in hunks)
    assert not any("压测" in h.new for h in hunks if h.op != "removed")  # 插入一行不会让后面的行都算变化
    assert "B1：=SUM(1,2) → =SUM(1,2,3)" in text  # 公式按文本比较
    assert any(h.location == "工作表「风险」" for h in hunks)


def test_csv_with_key_and_reordering():
    old = "编号,任务,状态\n1,联调,进行中\n2,压测,未开始\n3,文档,完成\n"
    new = "编号,任务,状态\n3,文档,完成\n2,压测,进行中\n1,联调,完成\n4,回归,未开始\n"
    hunks, flags, notes = diff_csv(old, new)
    assert sorted(ops(hunks)) == [("added", "编号=4"), ("modified", "编号=1"), ("modified", "编号=2")]
    assert "状态：进行中 → 完成" in next(h.new for h in hunks if h.location == "编号=1")
    assert "按主键列「编号」比较" in notes[0]

    shuffled = "编号,任务,状态\n3,文档,完成\n1,联调,进行中\n2,压测,未开始\n"
    hunks, flags, _ = diff_csv(old, shuffled)
    assert hunks == [] and flags == ["reordered"]


def test_csv_without_key_ignores_pure_reordering():
    old = "任务,状态\n联调,进行中\n联调,进行中\n压测,未开始\n"
    hunks, flags, notes = diff_csv(old, "任务,状态\n压测,未开始\n联调,进行中\n联调,进行中\n")
    assert hunks == [] and flags == ["reordered"] and "没有找到可靠的主键" in notes[0]
    hunks, _, _ = diff_csv(old, old + "回归,未开始\n")
    assert ops(hunks) == [("added", " 第 5 行")]


def test_diff_is_cached_and_recomputable(env):
    f = track(env, "a.md", "# A\n一")
    first = update(env, f, "a.md", "# A\n一\n二")
    snaps = env.svc.snapshots(f.id)
    differ = Differ(env.svc)
    cached = differ.diff(snaps[0], snaps[1])
    fresh = differ.diff(snaps[0], snaps[1], use_cache=False)
    assert first.to_dict() == cached.to_dict() == fresh.to_dict()
    assert (cached.from_snapshot_id, cached.to_snapshot_id) == (snaps[0].id, snaps[1].id)


def test_text_of_reparses_when_material_deleted(env):
    f = track(env, "a.md", "# A\n一")
    snap = env.svc.snapshots(f.id)[0]
    env.store.delete(snap.material_id)
    text, markdown = Differ(env.svc).text_of(snap)
    assert text == "# A\n一" and markdown


def test_parse_failure_is_flagged(env):
    f = track(env, "a.docx", docx_bytes(["正文"]))
    result = update(env, f, "a.docx", b"broken")
    assert result.hunks == [] and "parse_failed" in result.flags


def test_large_diff_is_truncated(env, monkeypatch):
    monkeypatch.setattr(config, "TRACKING_MAX_HUNKS", 5)
    f = track(env, "a.txt", "开头")
    result = update(env, f, "a.txt", "开头\n" + "\n\n".join(f"新段落 {i}" for i in range(20)))
    assert len(result.hunks) <= 5 and result.truncated + len(result.hunks) >= 1


def test_cli_diff(env, capsys):
    f = track(env, "a.md", "# A\n一")
    env.write("a.md", "# A\n一\n二")
    env.svc.scan()
    assert track_cli.main(["diff", str(f.id)], service=env.svc) == 0
    out = capsys.readouterr().out
    assert "v1 → v2" in out and "新增［A］：二" in out
