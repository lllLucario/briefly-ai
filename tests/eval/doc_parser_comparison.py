"""受控Word结构样本的真实.doc对照实验，不读取resources或调用模型。"""
import argparse
from datetime import datetime
import hashlib
import json
import math
import os
from pathlib import Path
import statistics
import subprocess
import threading
import time
from zoneinfo import ZoneInfo

from docx import Document
import psutil

from briefly import config
from briefly.legacy_word import converter_path
from briefly.loaders import load_document
from doc_parser_support import normalized,xhtml_markdown

REPO=Path(__file__).resolve().parents[2]


def save(path,data):
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def fixtures(folder):
    cases=[]
    def add(name,build,facts,headings=(),rows=(),absent=()):
        doc=Document()
        build(doc)
        path=folder/(name+".docx")
        doc.save(path)
        cases.append(dict(id=name,source=path.name,facts=facts,headings=list(headings),rows=list(rows),absent=list(absent)))
    def basic(d):
        d.add_heading("差旅住宿制度",1)
        d.add_paragraph("上海正式员工住宿上限650元，须在订房前审批；不适用于实习生和境外差旅。")
    add("01-条件",basic,["上海正式员工住宿上限650元","订房前审批","不适用于实习生和境外差旅"],["差旅住宿制度"])
    def hierarchy(d):
        for level,title,value in [(1,"产品2.0操作手册","适用版本2.0"),(2,"上传前置条件","须登录并有编辑权限"),(3,"E403处理","向管理员申请编辑权限，不得改名绕过")]:
            d.add_heading(title,level)
            d.add_paragraph(value)
    add("02-层级",hierarchy,["适用版本2.0","须登录并有编辑权限","不得改名绕过"],["产品2.0操作手册","上传前置条件","E403处理"])
    def merged(d):
        d.add_heading("报销额度表",1)
        t=d.add_table(rows=4,cols=3)
        data=[("项目","对象","额度"),("上海","正式员工","650元"),("杭州","","480元"),("合住","按房间核算","")]
        for r,row in enumerate(data):
            for c,value in enumerate(row): t.cell(r,c).text=value
        t.cell(1,1).merge(t.cell(2,1))
        t.cell(3,1).merge(t.cell(3,2))
    add("03-合并表格",merged,["650元","480元","按房间核算"],["报销额度表"],["| 杭州 | 正式员工 | 480元 |","| 合住 | 按房间核算 |  |"])
    def lists(d):
        d.add_heading("故障恢复步骤",1)
        for p in ["先核对任务状态","再确认是否执行成功","结果未知时不要自动重试"]:d.add_paragraph(p,style="List Number")
        d.add_paragraph("保留异常日志",style="List Bullet")
    add("04-列表",lists,["先核对任务状态","再确认是否执行成功","结果未知时不要自动重试","保留异常日志"],["故障恢复步骤"])
    def header(d):
        d.sections[0].header.paragraphs[0].text="页眉干扰：历史额度99999元"
        d.sections[0].footer.paragraphs[0].text="页脚干扰：不是业务事实"
        d.add_heading("有效制度",1)
        d.add_paragraph("本制度生效日期为2026-10-01，旧版不再适用。")
    add("05-页眉页脚",header,["生效日期为2026-10-01","旧版不再适用"],["有效制度"],absent=["历史额度99999元","不是业务事实"])
    def symbols(d):
        d.add_heading("字符与字段",1)
        d.add_paragraph('字段 rate_50%、C++、API_v2、20 MB、报价“待定”；预算为8万元，不代表已经支出。')
        d.add_paragraph("# 这行是正文，不是章节")
    add("06-符号",symbols,["rate_50%","C++","API_v2","20 MB","预算为8万元，不代表已经支出","# 这行是正文，不是章节"],["字符与字段"])
    def long(d):
        d.add_heading("长篇项目说明",1)
        for i in range(120):
            d.add_heading(f"第{i+1}节 事项",2)
            d.add_paragraph(f"事项{i+1}负责人为虚构人员甲，状态为待确认；不允许推断为完成。记录编号A{i:04d}。"*3)
        d.add_paragraph("文末关键条件：只有验收签字后才能上线。")
    add("07-长文",long,["记录编号A0000","记录编号A0119","文末关键条件：只有验收签字后才能上线"],["长篇项目说明","第120节 事项"])
    def tables(d):
        d.add_heading("产品参数清单",1)
        for i in range(8):
            t=d.add_table(rows=16,cols=3)
            for c,value in enumerate(["参数","值","条件"]):t.cell(0,c).text=value
            for r in range(1,16):
                t.cell(r,0).text=f"参数{i}-{r}"
                t.cell(r,1).text=f"{r*10}MB"
                t.cell(r,2).text="仅适用管理员"
        d.add_paragraph("清单末尾：不承诺普通账号具备同等权限。")
    add("08-多表格",tables,["参数0-1","参数7-15","150MB","不承诺普通账号具备同等权限"],["产品参数清单"],["| 参数7-15 | 150MB | 仅适用管理员 |"])
    # 复用项目既有自行构造的脚注/文本框XML样本，不读取外部资料。
    import importlib.util
    spec=importlib.util.spec_from_file_location("docx_fixture",REPO/"tests/test_docx.py")
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    (folder/"09-脚注文本框.docx").write_bytes(module._with_textbox_and_footnote())
    cases.append(dict(id="09-脚注文本框",source="09-脚注文本框.docx",facts=["文本框里的关键数据：预算30万","数据来源为2026年财务报表"],headings=[],rows=[],absent=["页眉：内部资料","页脚：第1页"]))
    return cases


def measured(action):
    stop=threading.Event()
    root=psutil.Process()
    start_rss=root.memory_info().rss
    peaks=dict(family_rss_bytes=start_rss,child_rss_bytes=0)
    def sample():
        while not stop.is_set():
            try:
                children=root.children(recursive=True)
                child=0
                for p in children:
                    try:child+=p.memory_info().rss
                    except psutil.Error:pass
                peaks['child_rss_bytes']=max(peaks['child_rss_bytes'],child)
                peaks['family_rss_bytes']=max(peaks['family_rss_bytes'],root.memory_info().rss+child)
            except psutil.Error:pass
            stop.wait(.02)
    thread=threading.Thread(target=sample,daemon=True);thread.start()
    started=time.perf_counter()
    try:
        value=action()
        return value,dict(seconds=time.perf_counter()-started,parent_rss_start=start_rss,**peaks)
    finally:
        stop.set();thread.join(timeout=2)


def quality(text,case):
    norm=normalized(text)
    return dict(facts=[normalized(f) in norm for f in case['facts']],
                headings=[any(line.lstrip().startswith('#') and normalized(h) in normalized(line) for line in text.splitlines()) for h in case['headings']],
                rows=[normalized(row) in norm for row in case['rows']],absent=[normalized(f) not in norm for f in case['absent']])


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('-n',type=int,default=3)
    args=parser.parse_args()
    config.DOC_PARSER='convert'  # 对照标签固定为转换链，不受本机默认解析路线变化影响。
    stamp=datetime.now(ZoneInfo('Australia/Sydney')).strftime('%Y%m%d-%H%M%S-%f')
    output=REPO/'.tmp/eval-runs'/('doc-comparison-'+stamp);output.mkdir(parents=True,exist_ok=False)
    docs=output/'samples';docs.mkdir()
    cases=fixtures(docs)
    save(output/'cases.json',cases)
    soffice=converter_path()
    options={'creationflags':subprocess.CREATE_NO_WINDOW} if os.name=='nt' else {}
    # 一次启动生成全部旧格式输入，样本生成时间不计入解析计时。
    subprocess.run([soffice,'-env:UserInstallation='+(output/'generation-profile').as_uri(),'--headless','--nologo','--norestore','--convert-to','doc:MS Word 97','--outdir',str(docs),*[str(docs/c['source']) for c in cases]],check=True,timeout=120,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,**options)
    java=next((REPO/'.tmp/doc-direct/runtime').rglob('java.exe'))
    jar=REPO/'.tmp/doc-direct/tika-app-3.3.3.jar'
    classes=REPO/'.tmp/doc-direct/classes'
    records=[]
    for repeat in range(1,args.n+1):
        for case in cases:
            source=docs/(case['id']+'.doc');payload=source.read_bytes()
            for method in ('conversion','direct_office') if repeat%2 else ('direct_office','conversion'):
                target=output/(f"{repeat}-{case['id']}-{method}");target.mkdir()
                def run():
                    if method=='conversion':
                        doc=load_document(source.name,payload)
                        return doc.text,dict(doc.report)
                    completed=subprocess.run([str(java),'-Xmx256m','-XX:ActiveProcessorCount=4','-Dfile.encoding=UTF-8','-Djava.awt.headless=true','-cp',str(jar)+os.pathsep+str(classes),'DocProbe',str(target),str(source)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=90,**options)
                    if completed.returncode:raise RuntimeError('直接解析失败')
                    text,warnings=xhtml_markdown((target/'0.xhtml').read_bytes())
                    return text,dict(warnings=warnings)
                try:
                    (text,report),stats=measured(run)
                    (target/'output.md').write_text(text,encoding='utf-8')
                    records.append(dict(case=case['id'],method=method,repeat=repeat,quality=quality(text,case),
                        stats=stats,report=report,input_bytes=len(payload),output_bytes=len(text.encode('utf-8')),error=None))
                except Exception as exc:
                    records.append(dict(case=case['id'],method=method,repeat=repeat,error=type(exc).__name__+': '+str(exc)))
                save(output/'results.json',records)
                print(repeat,case['id'],method,'error' if records[-1]['error'] else round(records[-1]['stats']['seconds'],3),flush=True)
    # 单JVM处理9文档，单独反映可复用进程的批量潜力，不与单文件计时混为一谈。
    batch=output/'direct-batch';batch.mkdir()
    _,batchstats=measured(lambda:subprocess.run([str(java),'-Xmx256m','-XX:ActiveProcessorCount=4','-Dfile.encoding=UTF-8','-cp',str(jar)+os.pathsep+str(classes),'DocProbe',str(batch),*[str(docs/(c['id']+'.doc')) for c in cases]],check=True,timeout=90,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,**options))
    save(output/'batch.json',dict(stats=batchstats,documents=len(cases),parse_times=(batch/'times.tsv').read_text(encoding='utf-8')))
    summary=[]
    for method in ('conversion','direct_office'):
        good=[r for r in records if r['method']==method and r['error'] is None]
        timings=sorted(r['stats']['seconds'] for r in good)
        summary.append(dict(method=method,runs=len(good),failed=sum(r['method']==method and r['error'] is not None for r in records),
            checks={k:dict(passed=sum(sum(r['quality'][k]) for r in good),total=sum(len(r['quality'][k]) for r in good)) for k in ('facts','headings','rows','absent')},
            p50_s=statistics.median(timings),p95_s=timings[math.ceil(.95*len(timings))-1],
            peak_child_rss=max(r['stats']['child_rss_bytes'] for r in good),peak_family_rss=max(r['stats']['family_rss_bytes'] for r in good)))
    save(output/'summary.json',summary)
    save(output/'manifest.json',dict(repeats=args.n,resources_read=False,external_model_calls=0,
        input_sha256={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in docs.glob('*.doc')},
        sources={str(p.relative_to(REPO)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [Path(__file__),Path(__file__).with_name('doc_parser_support.py'),REPO/'tests/eval/java/DocProbe.java']},
        method='Same-process Python host; Java starts per file, LibreOffice starts per file; alternated order; summed family RSS sampled at20ms, shared pages may be counted more than once',
        note='All DOC inputs generated by LibreOffice from controlled DOCX. Not a representative sample of all Word-authored DOC files.'))
    print(output.relative_to(REPO).as_posix(),flush=True)


if __name__=='__main__':main()
