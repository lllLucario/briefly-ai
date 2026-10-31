"""手动本地入库验证：真实解析/离线Embedding/SQLite索引；不调用生成模型。

输出仅保留状态、哈希与计数，不输出业务正文。原件需通过--document显式指定。
"""
import argparse
from io import BytesIO
import json
from pathlib import Path
import sys
from time import perf_counter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from briefly.db import create_db_engine
from briefly.knowledge.index import LocalIndex
from briefly.knowledge.ingestion import ImportService
from briefly.knowledge.evidence import restricted_answer
from briefly.knowledge.service import KnowledgeService
from briefly.llm import create_local_embeddings


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--model-dir", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--document", type=Path)
    args = parser.parse_args()
    root = Path(args.output).resolve()
    root.mkdir(parents=True, exist_ok=False)
    encoder = create_local_embeddings(args.model_dir)
    engine = create_db_engine("sqlite:///" + (root / "app.db").as_posix())
    report = dict(generation_model_calls=0, embedding_model=encoder.identity, files=[])
    try:
        svc = ImportService(engine, root, LocalIndex(root / "index.sqlite"), lambda: encoder)
        kb = svc.knowledge.create("原创虚构入库验证")
        from docx import Document
        from test_pdf import make_pdf
        doc = Document()
        doc.add_heading("虚构采购制度", 1)
        doc.add_paragraph("试验企业的采购申请须在下单前提交，超过五万元由财务复核。")
        table = doc.add_table(rows=2, cols=2)
        for cells, values in zip(table.rows, [("事项", "要求"), ("报销", "五日内提交")]):
            for cell, value in zip(cells.cells, values):
                cell.text = value
        stream = BytesIO()
        doc.save(stream)
        files = [("original.md", "# 原创虚构手册\n".encode() + "导出操作前应先保存当前草稿。".encode() * 100),
            ("original.docx", stream.getvalue()),
            ("original.pdf", make_pdf([[(10, 700, "Original fictional manual: save the draft before exporting.")]])),
            ("broken.docx", b"broken")]
        if args.document:
            files.append((args.document.name, args.document.read_bytes()))
        for number, (name, data) in enumerate(files):
            started = perf_counter()
            event = dict(filename=name)
            job = None
            try:
                job, _ = svc.enqueue(kb.id, name, data, request_key=f"validation:{number}")
                preview = svc.parse(kb.id, job.id)
                event["chunks"] = len(preview.preview["chunks"])
                ready = svc.confirm(kb.id, job.id, preview.preview["digest"])
                version = svc.knowledge.versions(kb.id, next(d.id for d in svc.knowledge.documents(kb.id) if d.source_key == job.source_key))[1][-1]
                item = svc.knowledge.chunk(kb.id, version.document_id, version.id, 1)
                event.update(status=ready.status, content_hash=version.content_hash, constraints=item["constraints"])
                if item["constraints"]:
                    event["answer_status"] = restricted_answer([item])["status"]
            except ValueError as exc:
                event.update(status=svc.get(kb.id, job.id).status if job else "rejected", error=str(exc))
            event["seconds"] = round(perf_counter() - started, 3)
            report["files"].append(event)
        report["scope_before_restart"] = svc.lifecycle.validate_scope([kb.id], encoder)["status"]
        report["foreign_external_api_calls"] = 0
    finally:
        engine.dispose()
    engine = create_db_engine("sqlite:///" + (root / "app.db").as_posix())
    try:
        svc = ImportService(engine, root, LocalIndex(root / "index.sqlite"), lambda: encoder)
        report["scope_after_restart"] = svc.lifecycle.validate_scope([kb.id], encoder)["status"]
        report["job_count_after_restart"] = len(svc.jobs(kb.id))
    finally:
        engine.dispose()
    (root / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"files": report["files"], "scope_after_restart": report["scope_after_restart"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
