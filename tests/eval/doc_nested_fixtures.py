"""原创嵌套DOC回归：父格归属、子表先出现、三级表与合并单元格。"""
import argparse
import json
from pathlib import Path
import subprocess
import os

from docx import Document
from briefly.direct_word import parse_docs
from briefly.legacy_word import converter_path
from doc_parser_comparison import save


def child(cell,rows,first=False):
    table=cell.add_table(rows=len(rows),cols=len(rows[0]))
    for r,row in enumerate(rows):
        for c,value in enumerate(row):table.cell(r,c).text=value
    if first:
        first_paragraph=cell.paragraphs[0]._p
        first_paragraph.getparent().remove(first_paragraph)
    return table


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('folder',type=Path);args=parser.parse_args()
    folder=args.folder.resolve();folder.mkdir(exist_ok=True)
    doc=Document();doc.add_heading('虚构嵌套差旅制度',1)
    table=doc.add_table(rows=3,cols=2)
    table.cell(0,0).text='城市';table.cell(0,1).text='规则'
    table.cell(1,0).text='上海';table.cell(1,1).text='须事前审批'
    child(table.cell(1,1),[['对象','额度'],['正式员工','650元'],['实习生','350元']])
    table.cell(2,0).text='杭州'
    child(table.cell(2,1),[['对象','额度'],['正式员工','480元'],['实习生','300元']],first=True)
    doc.add_paragraph('文末校验：END-NESTED-7391');doc.save(folder/'01-父格归属.docx')
    doc=Document();doc.add_heading('虚构三级嵌套表',1)
    table=doc.add_table(rows=2,cols=2)
    inner=child(table.cell(0,0),[['分类','规则'],['住宿','']],first=True)
    child(inner.cell(1,1),[['类型','额度'],['单人','650元'],['合住','按房间核算']],first=True)
    table.cell(0,1).text='审批负责人张三';table.cell(1,1).text='审批负责人李四'
    table.cell(0,0).merge(table.cell(1,0))
    doc.add_paragraph('文末校验：END-NESTED-7391');doc.save(folder/'02-三级合并.docx')
    sources=sorted(folder.glob('*.docx'))
    subprocess.run([converter_path(),'-env:UserInstallation='+(folder/'generation-profile').as_uri(),'--headless','--nologo','--norestore','--convert-to','doc:MS Word 97','--outdir',str(folder),*map(str,sources)],
        check=True,timeout=60,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
    records=[]
    for repeat in range(1,4):
        docs=parse_docs([p.with_suffix('.doc').read_bytes() for p in sources])
        for i,item in enumerate(docs):
            assert not isinstance(item,Exception),item
            text,report=item
            assert report['conversion'] is None and 'END-NESTED-7391' in text
            assert not report.get('unverified_table_segments'),report
            (folder/f'{i+1}-{repeat}.md').write_text(text,encoding='utf-8')
            if i==0:
                assert '| 上海 | 须事前审批 〔子表 2〕 |' in text
                assert '| 杭州 | 〔子表 3〕 |' in text
                assert '子表 2（父表 1，第2行，第2列）' in text
                assert '子表 3（父表 1，第3行，第2列）' in text
                assert '| 正式员工 | 650元 |' in text and '| 正式员工 | 480元 |' in text
                assert report['tables']==3 and report['nested_tables']==2
            else:
                assert '父表 1，第1至2行，第1列' in text
                assert '子表 3（父表 2，第2行，第2列）' in text
                assert '| 单人 | 650元 |' in text and '| 合住 | 按房间核算 |' in text
                assert '审批负责人张三' in text and '审批负责人李四' in text
                assert report['tables']==3 and report['nested_tables']==2
            records.append(dict(file=sources[i].name,repeat=repeat,report=report,passed=True))
    save(folder/'results.json',records)
    print('Nested DOC fixtures: 6 parses passed')


if __name__=='__main__':main()
