"""受控多格式解析资源测量；不读取用户资料、不调用模型。手动执行。"""
import argparse
from datetime import datetime
import hashlib
from importlib import metadata
import json
import os
from pathlib import Path
import statistics
import subprocess
import sys
import threading
import time
from zoneinfo import ZoneInfo

import psutil

REPO = Path(__file__).resolve().parents[2]


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def measure(action):
    process = psutil.Process()
    start_rss = process.memory_info().rss
    peak = dict(parent_rss=start_rss, children_rss=0)
    stop = threading.Event()
    def sample():
        while not stop.is_set():
            try:
                peak['parent_rss'] = max(peak['parent_rss'], process.memory_info().rss)
                children = 0
                for child in process.children(recursive=True):
                    try:
                        children += child.memory_info().rss
                    except psutil.Error:
                        pass
                peak['children_rss'] = max(peak['children_rss'], children)
            except psutil.Error:
                pass
            stop.wait(.02)
    thread = threading.Thread(target=sample, daemon=True)
    thread.start()
    start = time.perf_counter()
    try:
        result = action()
        return result, dict(seconds=time.perf_counter()-start, start_rss=start_rss, **peak)
    finally:
        stop.set()
        thread.join()


def worker(source, output):
    # 生成样本的依赖与API不导入本进程，避免把解析首次导入提前预热。
    payload = source.read_bytes()
    def imports():
        from briefly.loaders import load_document
        return load_document
    loader, import_stats = measure(imports)
    records = []
    for repeat in range(4):
        doc, stats = measure(lambda: loader(source.name, payload))
        assert 'PROBE-END-7391' in doc.text, source.name
        records.append(dict(repeat=repeat, stats=stats, chars=len(doc.text), report=doc.report))
    save(output, dict(file=source.name, bytes=len(payload), loader_import=import_stats, records=records,
                      java_loaded_in_parent=any(name.startswith('jpype') for name in sys.modules)))


def pdf_bytes(pages):
    # 自行构造带文字的PDF；不需要额外PDF生成依赖或字体下载。
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'', b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    kids = []
    for page in range(pages):
        index = len(objects)+1
        kids.append(f'{index} 0 R')
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> /Contents {index+1} 0 R >>'.encode())
        section = chr(65+page//26)+chr(65+page%26)
        lines = [f'Section {section} Company Manual', *[f'Section {section} Item {i}: staff allowance 650 requires approval before booking.' for i in range(35)], f'Section {section} PROBE-END-7391']
        content = ('BT /F1 10 Tf 40 800 Td 17 TL ' + ' '.join(f'({line}) Tj T*' for line in lines) + ' ET').encode()
        objects.append(f'<< /Length {len(content)} >>\nstream\n'.encode()+content+b'\nendstream')
    objects[1] = f'<< /Type /Pages /Kids [{" ".join(kids)}] /Count {pages} >>'.encode()
    result = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f'{index} 0 obj\n'.encode()+obj+b'\nendobj\n')
    xref = len(result)
    result.extend(f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode())
    result.extend(b''.join(f'{offset:010} 00000 n \n'.encode() for offset in offsets[1:]))
    result.extend(f'trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode())
    return bytes(result)


def fixtures(folder):
    from docx import Document
    from openpyxl import Workbook
    for size, count in [('small', 20), ('medium', 1000)]:
        doc = Document()
        doc.add_heading('虚构企业制度', 1)
        for i in range(count):
            doc.add_paragraph(f'第{i+1}条：正式员工在订房前须审批，上海住宿上限650元。' * 3)
        table = doc.add_table(rows=4, cols=3)
        for r, row in enumerate([['城市','对象','限额'],['上海','正式员工','650元'],['杭州','正式员工','480元'],['PROBE-END-7391','审核','完成']]):
            for c, value in enumerate(row):
                table.cell(r,c).text = value
        doc.save(folder/f'{size}.docx')
    for size, rows in [('small',100), ('medium',10000)]:
        book = Workbook()
        sheet = book.active
        sheet.title = '虚构项目资料'
        sheet.append([f'字段{i}' for i in range(10)])
        for row in range(rows):
            sheet.append([f'项目{row}', '负责人张三', 650, '待审批', *[f'值{row}-{i}' for i in range(6)]])
        sheet.cell(rows+2,1,'PROBE-END-7391')
        book.save(folder/f'{size}.xlsx')
        book.close()
    for size, pages in [('small',2),('medium',50)]:
        (folder/f'{size}.pdf').write_bytes(pdf_bytes(pages))
    (folder/'small.txt').write_text('虚构制度，650元，须审批。\n'*100+'PROBE-END-7391', encoding='utf-8')
    (folder/'small.csv').write_text('项目,负责人,额度\n'+'虚构项目,张三,650\n'*100+'PROBE-END-7391,审核,完成', encoding='utf-8')


def disk_inventory():
    from packaging.requirements import Requirement
    from packaging.utils import canonicalize_name
    def closure(root):
        pending = [Requirement(root)]
        seen = set()
        distributions = {}
        paths = set()
        while pending:
            requirement = pending.pop()
            name = canonicalize_name(requirement.name)
            key = (name, tuple(sorted(requirement.extras)))
            if key in seen:
                continue
            seen.add(key)
            dist = metadata.distribution(name)
            distributions[name] = dist.version
            for file in dist.files or []:
                path = Path(dist.locate_file(file)).resolve()
                if path.is_file():
                    paths.add(path)
            for raw in dist.requires or []:
                child = Requirement(raw)
                if child.marker is None or any(child.marker.evaluate({'extra': extra}) for extra in ['', *requirement.extras]):
                    pending.append(child)
        return dict(root=root, bytes=sum(path.stat().st_size for path in paths), files=len(paths), distributions=distributions,
                    note='Installed distribution RECORD files and active dependency closure; shared dependencies overlap, no sum across formats; excludes caches/unlisted pyc/Python runtime')
    return [closure(root) for root in ['docling-slim[format-docx]','openpyxl','pypdf']]


def concurrency(folder, output):
    from concurrent.futures import ThreadPoolExecutor
    from fastapi.testclient import TestClient
    from briefly import config, direct_word, loaders
    from briefly.api import create_app
    from briefly.db import create_db_engine
    from briefly.materials import MaterialStore
    source = REPO/'files/eval/w3/doc-comparison/samples/03-合并表格.doc'
    config.DOC_PARSER = 'direct'
    started, done = threading.Event(), threading.Event()
    original = direct_word._run
    def observed(*args, **kwargs):
        started.set()
        try:
            return original(*args, **kwargs)
        finally:
            done.set()
    def forbidden(*args, **kwargs):
        raise AssertionError('禁止模型调用')
    engine = create_db_engine('sqlite:///'+str((output/'concurrency.sqlite3').resolve()).replace('\\','/'))
    records = []
    try:
        with TestClient(create_app(MaterialStore(engine), model_factory=forbidden, agent_factory=forbidden)) as client:
            # 先确认普通上传在本轮没有DOC时的响应，再测同一服务并发DOC。
            def request(name, method, path, **kwargs):
                start = time.perf_counter()
                result = client.request(method, path, **kwargs)
                elapsed = time.perf_counter()-start
                assert result.status_code in (200,201), (name,result.text)
                return dict(name=name, seconds=elapsed, status=result.status_code, doc_still_running=started.is_set() and not done.is_set())
            for name in ['small.xlsx','small.docx','small.pdf']:
                records.append(request('baseline-'+name,'POST','/materials/upload',files={'file':(name,(folder/name).read_bytes())}))
            # 保留解析器预热，但删除基线材料，确保并发上传覆盖实际写入而非内容去重捷径。
            for material in client.get('/materials').json():
                assert client.delete('/materials/'+str(material['id'])).status_code == 204
            direct_word._run = observed
            with ThreadPoolExecutor(max_workers=2) as pool:
                future = pool.submit(request,'parallel-doc','POST','/materials/upload',files={'file':('probe.doc',source.read_bytes())})
                assert started.wait(20), 'DOC未启动'
                for _ in range(3):
                    records.append(request('parallel-list','GET','/materials'))
                for name in ['small.xlsx','small.docx','small.pdf']:
                    records.append(request('parallel-'+name,'POST','/materials/upload',files={'file':(name,(folder/name).read_bytes())}))
                records.append(future.result(timeout=90))
    finally:
        direct_word._run = original
        engine.dispose()
    save(output/'concurrency.json',dict(records=records, method='Real Java parser, FastAPI TestClient/threadpool, isolated SQLite, warm non-DOC parsers. In-process API, not network latency or saturation test.'))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--worker', type=Path)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--resume', type=Path)
    args = parser.parse_args()
    if args.worker:
        worker(args.worker,args.output)
        return
    output = args.resume or REPO/'.tmp/eval-runs'/('parser-resources-'+datetime.now(ZoneInfo('Australia/Sydney')).strftime('%Y%m%d-%H%M%S'))
    output = output.resolve()
    output.mkdir(exist_ok=bool(args.resume))
    folder = output/'samples'
    if not args.resume:
        folder.mkdir()
        fixtures(folder)
    print(output.relative_to(REPO), flush=True)
    results = []
    for source in sorted(folder.iterdir()):
        for repeat in range(3):
            target = output/f'{source.name}-{repeat}.json'
            if target.exists():
                record = json.loads(target.read_text(encoding='utf-8'))
                record['fresh_process'] = repeat
                results.append(record)
                continue
            started = time.perf_counter()
            completed = subprocess.run([sys.executable,str(Path(__file__)), '--worker',str(source),'--output',str(target)],
                stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=180,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
            (output/f'{source.name}-{repeat}.stderr.txt').write_bytes(completed.stderr)
            assert completed.returncode == 0, completed.stderr.decode(errors='replace')
            record = json.loads(target.read_text(encoding='utf-8'))
            record['process_wall_s'] = time.perf_counter()-started
            record['fresh_process'] = repeat
            results.append(record)
        print(source.name, 'done', flush=True)
        save(output/'results.json',results)
    summaries = []
    for name in sorted({r['file'] for r in results}):
        group = [r for r in results if r['file']==name]
        summaries.append(dict(file=name, input_bytes=group[0]['bytes'], output_chars=group[0]['records'][0]['chars'],
            first_p50_s=statistics.median(r['records'][0]['stats']['seconds'] for r in group),
            warm_p50_s=statistics.median(item['stats']['seconds'] for r in group for item in r['records'][1:]),
            first_peak_rss=max(r['records'][0]['stats']['parent_rss'] for r in group),
            first_rss_start=statistics.median(r['records'][0]['stats']['start_rss'] for r in group),
            max_children_rss=max(item['stats']['children_rss'] for r in group for item in r['records']),
            process_wall_p50_s=statistics.median(r['process_wall_s'] for r in group if 'process_wall_s' in r) if any('process_wall_s' in r for r in group) else None))
    save(output/'summary.json', summaries)
    save(output/'disk.json', disk_inventory())
    concurrency(folder,output)
    save(output/'manifest.json',dict(script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(), python=sys.version,
        cpu=psutil.cpu_count(), ram=psutil.virtual_memory().total, resources_read=False, external_model_calls=0,
        method='Three fresh Python processes per input, each first parse plus three warm parses; 20ms RSS sampling. Fixture generation excluded. OS disk cache not flushed. Script imports psutil before measurement. File read excluded from parse timer. Not answer quality or production load test.',
        samples={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.iterdir()}))
    print('ALL DONE', output.relative_to(REPO), flush=True)


if __name__ == '__main__':
    main()
