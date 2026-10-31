"""端到端问答评测：虚构企业资料（Word＋PDF）→ 本地入库 → 助理完整链路回答 → 规则判分。

手动运行，会把问题和检索到的虚构资料片段发给已配置的生成 API（需 --allow-external-fictional）。
资料与题目在 tests/eval/fixtures/w3_e2e/：documents.json 为虚构资料，cases.jsonl 为独立出题的题集。

用法（仓库根目录）：
    python tests/eval/w3_e2e_qa.py build .tmp/e2e-kb             # 生成 docx/PDF 并本地入库，不外发
    python tests/eval/w3_e2e_qa.py run .tmp/e2e-kb --allow-external-fictional -n 3 [--split dev] [--path assistant|qa] [-k R01 C02]

--path assistant（默认）走网页助理同一链路：意图漏斗 → 知识捷径或智能体调用 search_knowledge → 回答；
--path qa 直接调用知识问答接口同一函数（KnowledgeQA.answer）。
结果写入 files/eval/w3/e2e/<时间>-<path>/：results.json（逐次明细）和 report.md（汇总）。
holdout 题只看汇总，不看明细调参。
"""
import argparse
from collections import Counter
from datetime import datetime
import json
from pathlib import Path
import subprocess
import sys
import tempfile
from time import perf_counter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # 让 tests.eval 可导入
from tests.eval.w3_e2e_scoring import score, summarize  # noqa: E402

from briefly import config  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures" / "w3_e2e"
KB_NAME = "端到端评测虚构资料库"


def load_documents() -> list[dict]:
    return json.loads((FIXTURES / "documents.json").read_text(encoding="utf-8"))["documents"]


def load_cases(split: str, keys: list[str]) -> list[dict]:
    cases = [json.loads(line) for line in (FIXTURES / "cases.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    if keys:
        return [c for c in cases if c["id"] in keys]
    return [c for c in cases if split == "all" or c["split"] == split]


def write_docx(doc: dict, path: Path) -> None:
    from docx import Document
    word = Document()
    word.core_properties.title = doc["title"]
    word.add_heading(doc["title"], level=0)
    for kind, value in doc["blocks"]:
        if kind in ("h1", "h2"):
            word.add_heading(value, level=1 if kind == "h1" else 2)
        elif kind == "p":
            word.add_paragraph(value)
        elif kind == "table":
            table = word.add_table(rows=len(value), cols=len(value[0]))
            table.style = "Table Grid"
            for r, row in enumerate(value):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    word.save(path)


def to_pdf(source: Path, out_dir: Path) -> Path:
    from briefly.legacy_word import converter_path
    with tempfile.TemporaryDirectory() as profile:  # 独立配置目录，不受已打开的 LibreOffice 影响
        subprocess.run([converter_path(), f"-env:UserInstallation={Path(profile).as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(out_dir), str(source)],
                       check=True, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return out_dir / (source.stem + ".pdf")


def open_imports(root: Path, encoder):
    from briefly.db import create_db_engine
    from briefly.knowledge.index import LocalIndex
    from briefly.knowledge.ingestion import ImportService
    engine = create_db_engine("sqlite:///" + (root / "app.db").as_posix())
    return engine, ImportService(engine, root, LocalIndex(root / "index.sqlite"), lambda: encoder)


def build(root: Path) -> None:
    from briefly.llm import create_local_embeddings
    root.mkdir(parents=True, exist_ok=False)
    files = root / "files"
    files.mkdir()
    encoder = create_local_embeddings(config.EMBEDDING_DIR)
    engine, imports = open_imports(root, encoder)
    manifest = {}
    try:
        kb = imports.knowledge.create(KB_NAME)
        for doc in load_documents():
            path = files / f"{doc['title']}.docx"
            write_docx(doc, path)
            if doc["format"] == "pdf":
                pdf = to_pdf(path, files)
                path.unlink()
                path = pdf
            job, _ = imports.enqueue(kb.id, path.name, path.read_bytes(), request_key=doc["id"], source_key=doc["id"])
            preview = imports.parse(kb.id, job.id)
            ready = imports.confirm(kb.id, job.id, preview.preview["digest"])
            manifest[doc["id"]] = dict(file=path.name, format=doc["format"], status=ready.status,
                                       chunks=len(preview.preview["chunks"]))
            print(doc["id"], path.name, ready.status, len(preview.preview["chunks"]), flush=True)
        (root / "manifest.json").write_text(json.dumps(dict(kb_id=kb.id, documents=manifest), ensure_ascii=False, indent=1),
                                            encoding="utf-8")
    finally:
        engine.dispose()


def run(root: Path, args) -> None:
    from briefly.llm import create_local_embeddings

    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    documents = {d["id"]: d for d in load_documents()}
    cases = load_cases(args.split, args.k)
    encoder = create_local_embeddings(config.EMBEDDING_DIR)
    engine, imports = open_imports(root, encoder)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    try:
        kb_id = manifest["kb_id"]
        bases = imports.knowledge.list_bases()
        if [b.name for b in bases] != [KB_NAME]:
            raise SystemExit("不是本脚本建的虚构资料库，拒绝外发")
        imports.knowledge.set_egress(kb_id, True)  # 已核对为专用虚构资料库，按库授权外发
        rows, events = answer_cases(imports, encoder, kb_id, cases, lambda c: case_format(c, documents), args, stamp)
    finally:
        engine.dispose()
    write_report(Path(args.output) / f"{stamp}-{args.path}", args, root, manifest["documents"], rows, events, stamp)


def answer_cases(imports, encoder, kb_id, cases, doc_format_of, args, stamp):
    """逐题按 args.path 指定的链路回答 args.n 次并判分，返回（汇总行, 逐次明细）。外发授权由调用方核对。"""
    from briefly.agent import Agent
    from briefly.intent import IntentFunnel
    from briefly.knowledge.assistant import KnowledgeAssistant
    from briefly.knowledge.retrieval import RetrievalService
    from briefly.llm import create_chat_model

    rows, events = [], []
    titles, corpus = library(imports, kb_id)
    retrieval = RetrievalService(imports)
    model_factory = lambda: create_chat_model(max_retries=1, timeout=90)
    knowledge = KnowledgeAssistant(retrieval, model_factory)
    if args.path == "assistant":
        agent = Agent(model_factory(), store=imports.store)
        agent.enable_knowledge(knowledge, IntentFunnel(lambda: encoder, kb_probe=knowledge.relevance))
    for case in cases:
        evidence = local_evidence(retrieval, kb_id, case["question"])
        doc_format = doc_format_of(case)
        for repeat in range(args.n):
            started = perf_counter()
            event = dict(id=case["id"], category=case["category"], split=case["split"], repeat=repeat)
            try:
                if args.path == "assistant":
                    thread = f"e2e-{case['id']}-{repeat}-{stamp}"
                    answer = agent.chat(case["question"], thread_id=thread, kb_id=kb_id)
                    history = agent.history_of(thread)
                    human = next(m for m in reversed(history) if m.type == "human")
                    reply = history[-1]
                    meta = reply.response_metadata.get("knowledge") or {}
                    intent = human.additional_kwargs.get("intent") or {}
                    status = meta.get("status")
                    event.update(intent=intent.get("intent"), layer=intent.get("layer"), resolved=intent.get("resolved"),
                                 tools=[c["name"] for m in history if m.type == "ai" for c in getattr(m, "tool_calls", [])],
                                 forced_search=any(m.type == "ai" and m.response_metadata.get("forced_search") for m in history))
                else:
                    result = knowledge.qa.answer(kb_id, case["question"])
                    answer, status = result["answer"], result["status"]
                event.update(status=status, answer=answer)
                scored = judge(case, answer, status, evidence, titles, corpus)
            except Exception as exc:  # 记录后继续，失败单独统计
                event.update(status="error", error=f"{type(exc).__name__}: {exc}"[:300])
                scored = dict(passed=False, problems=["error"], must=[], abstained=False, retrieval_hit=None,
                              unsupported_numbers=[], length=0, gold_withheld=None)
            event.update(scored, seconds=round(perf_counter() - started, 1))
            events.append(event)
            rows.append(dict(scored, category=case["category"], split=case["split"], doc_format=doc_format))
            print(case["id"], repeat, "PASS" if scored["passed"] else ",".join(scored["problems"]), flush=True)
    return rows, events


def write_report(out: Path, args, root: Path, documents: dict, rows, events, stamp) -> None:
    out.mkdir(parents=True, exist_ok=False)
    report = dict(path=args.path, n=args.n, split=args.split, model=config.MODEL, created=stamp, kb_root=str(root),
                  documents=documents,
                  summary={split: summarize([r for r in rows if split == "all" or r["split"] == split])
                           for split in ("all", "dev", "holdout")},
                  errors=sum(e["status"] == "error" for e in events))
    (out / "results.json").write_text(json.dumps(dict(report, events=events), ensure_ascii=False, indent=1), encoding="utf-8")
    (out / "report.md").write_text(render(report, events), encoding="utf-8")
    print(f"结果：{out}")


def library(imports, kb_id):
    """资料编号 → 入库标题（依据片段按标题对应资料），以及全部资料原文（核对数字出处）。"""
    titles, corpus = {}, []
    for item in imports.knowledge.documents(kb_id):
        material = imports.store.get(imports.knowledge.versions(kb_id, item.id)[1][-1].material_id, workspace=False)
        titles[item.source_key] = material.title
        corpus.append(material.content)
    return titles, corpus


def local_evidence(retrieval, kb_id, question):
    """按原问题检索（本地，不外发），参数与本轮依据包相同；保留质量限制剔除前的结果，
    这样“检索没找到”和“找到了但因结构未确认被剔除”能分开统计。"""
    result = retrieval.search(kb_id, question, include_document_context=True, context_budget_only=True,
                              context_tokens=config.KNOWLEDGE_CONTEXT_TOKENS)
    return [dict(title=e["source"]["title"], content=e["content"], withheld=bool(e.get("constraints")))
            for e in result["evidence"]]


def case_format(case, documents):
    formats = {documents[g["doc"]]["format"] for g in case["gold"]}
    return "pdf" if "pdf" in formats else "docx" if formats else "none"


def judge(case, answer, status, evidence, titles, corpus):
    scored = score(case, answer, status, evidence, titles, corpus)
    if scored["retrieval_hit"]:
        # 命中的依据是否全部带质量限制（不会交给模型）
        usable = [e for e in evidence if not e["withheld"]]
        scored["gold_withheld"] = not score(case, answer, status, usable, titles, corpus)["retrieval_hit"]
    else:
        scored["gold_withheld"] = None
    return scored


def rescore(results: Path) -> None:
    """不调用模型：按当前判分规则和本地检索结果重新判分已有运行，覆盖同目录的 results.json 与 report.md。"""
    from briefly.knowledge.retrieval import RetrievalService
    from briefly.llm import create_local_embeddings
    data = json.loads((results / "results.json").read_text(encoding="utf-8"))
    root = Path(data.get("kb_root", ".tmp/e2e-kb"))
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    documents = {d["id"]: d for d in load_documents()}
    cases = {c["id"]: c for c in load_cases("all", [])}
    engine, imports = open_imports(root, create_local_embeddings(config.EMBEDDING_DIR))
    try:
        titles, corpus = library(imports, manifest["kb_id"])
        retrieval = RetrievalService(imports)
        evidence = {key: local_evidence(retrieval, manifest["kb_id"], cases[key]["question"])
                    for key in {e["id"] for e in data["events"]}}
    finally:
        engine.dispose()
    rows = []
    for event in data["events"]:
        case = cases[event["id"]]
        if event["status"] != "error":
            event.update(judge(case, event["answer"], event["status"], evidence[event["id"]], titles, corpus))
        rows.append(dict({k: event[k] for k in ("passed", "problems", "retrieval_hit", "length")},
                         gold_withheld=event.get("gold_withheld"), category=case["category"], split=case["split"],
                         doc_format=case_format(case, documents)))
    data["summary"] = {split: summarize([r for r in rows if split == "all" or r["split"] == split])
                       for split in ("all", "dev", "holdout")}
    data["rescored"] = datetime.now().strftime("%Y%m%d-%H%M%S")
    (results / "results.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    (results / "report.md").write_text(render(data, data["events"]), encoding="utf-8")
    print(f"已重新判分：{results}")


def render(report: dict, events: list[dict]) -> str:
    lines = [f"# 端到端问答评测 {report['created']}", "",
             f"- 链路：{report['path']}；模型：{report['model']}；每题 {report['n']} 次；题集：{report['split']}；出错 {report['errors']} 次",
             "- 判分为规则核对（要点串、拒答写法、无出处数字、字数），不是人工或模型评审；失败明细需人工复核。",
             "- 检索命中按质量限制剔除前的本地检索结果计；“命中但依据被剔除”指命中的片段全部带结构未确认限制、不会交给模型。", ""]
    for split, groups in report["summary"].items():
        if not groups["all"]["runs"]:
            continue
        lines += [f"## {split}", "", "| 分组 | 次数 | 通过率 | 检索命中 | 命中但依据被剔除 | 检索失败 | 答非所问 | 幻觉 | 冗余 | 字数中位 |",
                  "|---|---|---|---|---|---|---|---|---|---|"]
        for name, g in groups.items():
            p = g["problems"]
            lines.append(f"| {name} | {g['runs']} | {g['pass_rate']} | {g['retrieval_hit']} | {g.get('gold_withheld')} | {p['retrieval_miss']} | "
                         f"{p['off_target']} | {p['hallucination']} | {p['verbose']} | {g['median_length']} |")
        lines.append("")
    failed = Counter(e["id"] for e in events if e["split"] == "dev" and not e["passed"])
    if failed:
        lines += ["## dev 未通过的题（次数）", "", "、".join(f"{k}×{v}" for k, v in sorted(failed.items())), ""]
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    b = sub.add_parser("build")
    b.add_argument("root", type=Path)
    r = sub.add_parser("run")
    r.add_argument("root", type=Path)
    r.add_argument("--allow-external-fictional", action="store_true")
    r.add_argument("--path", choices=["assistant", "qa"], default="assistant")
    r.add_argument("--split", choices=["dev", "holdout", "all"], default="all")
    r.add_argument("-n", type=int, default=3)
    r.add_argument("-k", nargs="*", default=[])
    r.add_argument("--output", default="files/eval/w3/e2e")
    s = sub.add_parser("rescore", help="不调用模型，按当前规则重新判分已有运行")
    s.add_argument("results", type=Path, help="某次运行的输出目录")
    args = parser.parse_args()
    if not config.EMBEDDING_DIR:
        parser.error("需要本地 Embedding 目录（.env 的 BRIEFLY_EMBEDDING_DIR）")
    if args.command == "build":
        build(args.root)
    elif args.command == "rescore":
        rescore(args.results)
    else:
        if not args.allow_external_fictional:
            parser.error("需显式允许把虚构问题与检索片段发往已配置的生成 API：--allow-external-fictional")
        run(args.root, args)


if __name__ == "__main__":
    main()
