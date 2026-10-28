"""用自行生成的中文Word验证真实本地转换，不读取用户资源，不调用模型。"""
import argparse
from datetime import datetime
import json
import os
from pathlib import Path
import subprocess
from zoneinfo import ZoneInfo

from docx import Document
from fastapi.testclient import TestClient

from briefly import config
from briefly.api import create_app
from briefly.db import create_db_engine
from briefly.legacy_word import converter_path
from briefly.loaders import load_document, load_text
from briefly.materials import MaterialStore


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--converter", help="本地soffice程序；未指定时用项目配置")
    args=parser.parse_args()
    if args.converter:
        config.LIBREOFFICE_PATH=args.converter
    executable=converter_path()
    repo=Path(__file__).resolve().parents[2]
    stamp=datetime.now(ZoneInfo("Australia/Sydney")).strftime("%Y%m%d-%H%M%S-%f")
    folder=repo/"files"/"eval"/("doc-walkthrough-"+stamp)
    folder.mkdir(parents=True,exist_ok=False)
    samples=folder/"中文路径 带空格（测试）"
    samples.mkdir()
    source=samples/"虚构制度.docx"
    doc=Document()
    doc.add_heading("虚构住宿制度",level=1)
    doc.add_paragraph("上海住宿上限650元，须在订房前审批，不适用于实习生。")
    table=doc.add_table(rows=2,cols=2)
    for (r,c),text in { (0,0):"城市",(0,1):"上限",(1,0):"杭州",(1,1):"480元" }.items():
        table.cell(r,c).text=text
    doc.save(source)
    options={"creationflags":subprocess.CREATE_NO_WINDOW} if os.name=="nt" else {}
    completed=subprocess.run([executable,"-env:UserInstallation="+(folder/"生成配置").as_uri(),
        "--headless","--nologo","--nodefault","--norestore","--convert-to","doc:MS Word 97",
        "--outdir",str(samples),str(source)],stdin=subprocess.DEVNULL,stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,timeout=config.DOC_CONVERSION_TIMEOUT_S,**options)
    legacy=samples/"虚构制度.doc"
    if completed.returncode or not legacy.is_file():
        raise RuntimeError("无法生成旧版Word测试文档")
    parsed=load_document(legacy.name,legacy.read_bytes())
    checks={phrase:phrase in parsed.text for phrase in ["虚构住宿制度","650元","订房前审批","不适用于实习生","杭州","480元"]}
    assert all(checks.values())
    assert "650元" in load_text(legacy)
    def forbidden(*args,**kwargs):
        raise AssertionError("元数据上传流程不应调用模型")
    engine=create_db_engine("sqlite://")
    try:
        with TestClient(create_app(MaterialStore(engine),model_factory=forbidden,agent_factory=forbidden)) as client:
            response=client.post("/materials/upload",files={"file":(legacy.name,legacy.read_bytes())})
            assert response.status_code==201,response.text
            record=dict(checks=checks,report=parsed.report,upload=response.json(),resources_read=False,
                        external_model_calls=0,original_bytes=legacy.stat().st_size)
        (folder/"result.json").write_text(json.dumps(record,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        print(folder.relative_to(repo).as_posix())
    finally:
        engine.dispose()


if __name__=="__main__":
    main()
