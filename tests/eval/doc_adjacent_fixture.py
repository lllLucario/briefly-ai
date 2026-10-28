"""真实普通段落分隔两张独立表格：修复不能误合并。原创材料，本地转换。"""
import argparse,json,subprocess,os
from pathlib import Path
from docx import Document
from briefly.direct_word import parse_docs
from briefly.legacy_word import converter_path


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('folder',type=Path);folder=parser.parse_args().folder.resolve();folder.mkdir(exist_ok=True)
    doc=Document();doc.add_heading('相邻独立表格边界回归',1)
    for i in range(2):
        t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='表'+str(i+1);t.cell(0,1).text='额度';t.cell(1,0).text='对象'+str(i+1);t.cell(1,1).text=str((i+1)*100)+'元'
        if i==0:doc.add_paragraph('普通正文边界 ADJACENT-PROSE-6482：下一张表属于独立业务。')
    doc.add_paragraph('END-ADJACENT-6482');doc.save(folder/'相邻表格.docx')
    subprocess.run([converter_path(),'-env:UserInstallation='+(folder/'profile').as_uri(),'--headless','--nologo','--norestore','--convert-to','doc:MS Word 97','--outdir',str(folder),str(folder/'相邻表格.docx')],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=60,creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
    results=[]
    for repeat in range(1,4):
        item=parse_docs([(folder/'相邻表格.doc').read_bytes()])[0];assert isinstance(item,tuple),item
        text,report=item;assert report['tables']==2 and report['nested_tables']==0
        assert report['unverified_table_segments']==0 and report['restored_source_paragraphs']==0
        assert text.index('| 对象1 | 100元 |')<text.index('ADJACENT-PROSE-6482')<text.index('| 表2 | 额度 |')
        assert not any('ADJACENT-PROSE-6482' in line and line.startswith('|') for line in text.splitlines())
        (folder/f'{repeat}.md').write_text(text,encoding='utf-8');results.append(dict(repeat=repeat,passed=True,report=report))
    (folder/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('3 adjacent-table parses preserve the ordinary paragraph as a separate boundary')


if __name__=='__main__':main()
