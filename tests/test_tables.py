from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from briefly import config
from briefly.chunking import chunk_text, parse_sections
from briefly.loaders import load_document
from briefly.tables import to_markdown_table


def xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_cells_are_normalized():
    table = to_markdown_table(
        [
            ["任务", "工时", "日期", "备注"],
            ["支付|改造", 16.0, datetime(2026, 10, 20), "第一行\n第二行"],
            [None, None, None, None],  # 空行跳过
            ["压测", 8.5, datetime(2026, 10, 18, 9, 30), None],
        ]
    )
    assert table == (
        "| 任务 | 工时 | 日期 | 备注 |\n"
        "|---|---|---|---|\n"
        "| 支付\\|改造 | 16 | 2026-10-20 | 第一行 第二行 |\n"
        "| 压测 | 8.5 | 2026-10-18 09:30 |  |"
    )


def test_trailing_empty_columns_are_dropped():
    assert to_markdown_table([["a", "b", None, None], ["1", None, None, None]]) == "| a | b |\n|---|---|\n| 1 |  |"


def test_xlsx_sheets_become_sections():
    data = xlsx_bytes({"任务进度": [["任务", "负责人"], ["支付改造", "李经理"]], "会议记录": [["日期", "主题"], ["10-08", "评审"]]})
    doc = load_document("进度.xlsx", data)
    assert doc.markdown is True and doc.report["structure"] == "sheets"
    assert [s.title for s in parse_sections(doc.text, doc.markdown)] == ["任务进度", "会议记录"]
    assert "| 支付改造 | 李经理 |" in doc.text


def test_formula_without_cached_value_is_reported():
    data = xlsx_bytes({"工时": [["任务", "工时"], ["a", 1], ["b", 2], ["合计", "=SUM(B2:B3)"]]})
    doc = load_document("t.xlsx", data)
    assert "1 个公式单元格没有计算结果" in doc.report["warnings"][0]


def test_xlsx_merged_cells_follow_the_same_rule_as_docx():
    wb = Workbook()
    ws = wb.active
    for row in [["工时", None, "任务", "负责人"], ["计划", "实际", None, None], [8, 10, "支付", "李经理"], [4, None, "压测", None], ["备注：按计划", None, None, None]]:
        ws.append(row)
    ws.merge_cells("A1:B1")  # 表头横向：每列都填；下一行落在它下面的“计划/实际”是第二层表头
    ws.merge_cells("D3:D4")  # 纵向：向下填
    ws.merge_cells("A5:D5")  # 正文横向：只留第一格
    buf = BytesIO()
    wb.save(buf)
    lines = load_document("t.xlsx", buf.getvalue()).text.splitlines()
    assert lines[2] == "| 工时/计划 | 工时/实际 | 任务 | 负责人 |"
    assert lines[5] == "| 4 |  | 压测 | 李经理 |"
    assert lines[6] == "| 备注：按计划 |  |  |  |"


def test_xlsx_caption_rows_and_multi_level_header_become_one_header_line():
    wb = Workbook()
    ws = wb.active
    rows = [
        [None, "薪酬结构总表", None, None, None],
        [None, None, None, None, "单位:元"],
        [None, "薪酬等级", "月工资", None, "备注"],
        [None, None, "标准\n工资", "绩效 工资", None],
        [None, "24级", 200000, 50000, "按月发放"],
        [None, "薪酬等级", "月工资", None, "备注"],  # 分段子表重复的表头
        [None, None, "标准\n工资", "绩效 工资", None],
        [None, "23级", 140000, 50000, None],
    ]
    for row in rows:
        ws.append(row)
    for rng in ("B1:E1", "B3:B4", "C3:D3", "E3:E4", "B6:B7", "C6:D6", "E6:E7"):
        ws.merge_cells(rng)
    buf = BytesIO()
    wb.save(buf)
    text = load_document("t.xlsx", buf.getvalue()).text
    assert text == (
        "# Sheet\n\n薪酬结构总表\n单位:元\n\n"
        "| 薪酬等级 | 月工资/标准工资 | 月工资/绩效工资 | 备注 |\n"
        "|---|---|---|---|\n"
        "| 24级 | 200000 | 50000 | 按月发放 |\n"
        "| 23级 | 140000 | 50000 |  |"
    )  # 空的 A 列去掉；表名不再当表头；“标准\n工资”的换行去掉


def test_xlsx_schedule_columns_are_compressed_to_start_and_end():
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "任务"] + list(range(1, 11)))
    ws.append([1, "寻址"] + [None, "■", "■", "■", None, None, None, None, "■", None])
    ws.append([2, "装修"] + [None] * 10)
    ws.append(["特别注意：网络平稳过渡"] + [None] * 11)
    ws.append([None] * 12)
    ws.merge_cells("A4:L5")  # 横跨整表的说明框：只出现一次，不向下填成多行
    buf = BytesIO()
    wb.save(buf)
    assert load_document("t.xlsx", buf.getvalue()).text.splitlines()[2:] == [
        "| 序号 | 任务 | 排期（1～10） |",
        "|---|---|---|",
        "| 1 | 寻址 | 2～4、9 |",
        "| 2 | 装修 |  |",
        "| 特别注意：网络平稳过渡 |  |  |",
    ]


def test_xlsx_month_ticks_under_a_group_header_are_compressed():
    wb = Workbook()
    ws = wb.active
    ws.append(["工作内容", "计划完成时间"] + [None] * 11 + ["输出"])
    ws.append([None] + [f"{m}月" for m in range(1, 13)] + [None])
    ws.append(["风险自查", "√"] + [None] * 5 + ["√"] + [None] * 5 + ["《自查表》"])
    for rng in ("A1:A2", "B1:M1", "N1:N2"):
        ws.merge_cells(rng)
    buf = BytesIO()
    wb.save(buf)
    assert load_document("t.xlsx", buf.getvalue()).text.splitlines()[2:] == [
        "| 工作内容 | 计划完成时间（1月～12月） | 输出 |",
        "|---|---|---|",
        "| 风险自查 | 1月、7月 | 《自查表》 |",
    ]


def test_xlsx_numeric_month_columns_are_not_compressed():
    data = xlsx_bytes({"预算": [["部门"] + list(range(1, 13)), ["销售"] + [100 * i for i in range(1, 13)]]})
    text = load_document("t.xlsx", data).text
    assert "| 部门 | 1 | 2 |" in text and "| 销售 | 100 | 200 |" in text


def test_xlsx_percent_is_shown_as_in_excel():
    wb = Workbook()
    ws = wb.active
    ws.append(["指标", "完成率", "增长"])
    ws.append(["支付", 0.15, 0.1234])
    ws["B2"].number_format = "0%"
    ws["C2"].number_format = "0.00%"
    buf = BytesIO()
    wb.save(buf)
    assert "| 支付 | 15% | 12.34% |" in load_document("t.xlsx", buf.getvalue()).text


@pytest.mark.parametrize(
    "text",
    ["任务,负责人\n支付改造,李经理\n", "任务;负责人\n支付改造;李经理\n", "任务\t负责人\n支付改造\t李经理\n"],
)
def test_csv_delimiters(text):
    doc = load_document("a.csv", text.encode("utf-8"))
    assert doc.text == "| 任务 | 负责人 |\n|---|---|\n| 支付改造 | 李经理 |"


def test_gbk_csv_from_excel():
    doc = load_document("a.csv", "任务,负责人\n压测,小王\n".encode("gbk"))
    assert "| 压测 | 小王 |" in doc.text


def test_long_table_chunks_repeat_the_header(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE", 120)
    rows = [["任务", "负责人", "状态"]] + [[f"任务{i:02d}", "某人", "进行中"] for i in range(30)]
    text = "# 进度\n\n说明文字。\n\n" + to_markdown_table(rows)
    chunks = chunk_text(text, parse_sections(text, markdown=True))
    assert len(chunks) > 3
    header = "| 任务 | 负责人 | 状态 |\n|---|---|---|\n"
    assert all(c.content.count("| 任务 | 负责人 | 状态 |") == 1 for c in chunks)
    assert all(c.content.startswith(header) for c in chunks[1:])
    body = "".join(c.content.removeprefix(header) for c in chunks[1:])
    assert all(f"任务{i:02d}" in chunks[0].content + body for i in range(30))  # 每一行都在
